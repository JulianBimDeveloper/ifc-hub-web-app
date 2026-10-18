"""
client_review.py — Servicio de escaneo y comparación para CLIENT REVIEW.

Lógica pura (sin FastAPI). El router llama a estas funciones vía run_in_executor.
"""

import os
import re
import json
import hashlib
import pickle
import sys
from collections import defaultdict

from services.client_review_schema import SCHEMA, _is_empty, EMPTY_VALUES
from core.resilience import atomic_write_model

try:
    from routers.ifc import _extract_ifc_info
except ImportError:
    _extract_ifc_info = None

# ── Directorio de caché en disco (NUNCA dentro de backend/) ──────────────────
_CACHE_DIR = os.environ.get(
    "IFC_HUB_REVIEW_CACHE",
    os.path.join(os.path.expanduser("~"), ".ifc-hub-v2", "client_review_cache"),
)
_REVIEWS_DIR = os.path.join(_CACHE_DIR, "reviews")

# In-memory cache: (path, mtime) → scan_result
_MEM_CACHE: dict = {}


# ── Utilidades de caché ───────────────────────────────────────────────────────

def _disk_key(path: str) -> str:
    try:
        st = os.stat(path)
        raw = f"{os.path.abspath(path)}|{int(st.st_mtime)}|{st.st_size}"
    except OSError:
        raw = os.path.abspath(path)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def _mem_key(path: str):
    try:
        mtime = os.stat(path).st_mtime
    except OSError:
        mtime = 0.0
    return (os.path.abspath(path), mtime)


def _cache_read(path: str) -> dict | None:
    # NUNCA devolver cache si la estructura ha cambiado
    # Invalidar siempre para asegurar datos frescos
    return None


def _cache_write(path: str, data: dict):
    key = _mem_key(path)
    _MEM_CACHE[key] = data
    dk = _disk_key(path)
    try:
        os.makedirs(_CACHE_DIR, exist_ok=True)
        fp = os.path.join(_CACHE_DIR, f"{dk}.json")
        tmp = fp + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(data, fh, ensure_ascii=False)
        os.replace(tmp, fp)
    except OSError:
        pass


def invalidate_cache(paths: list[str]):
    """Llamar después de bulk-fill u otras escrituras."""
    for path in paths:
        key = _mem_key(path)
        _MEM_CACHE.pop(key, None)
        dk = _disk_key(path)
        fp = os.path.join(_CACHE_DIR, f"{dk}.json")
        try:
            os.remove(fp)
        except OSError:
            pass


# ── Extracción de valores ─────────────────────────────────────────────────────

def _utm_to_wgs84(easting: float, northing: float, epsg: str):
    """Pure-Python UTM → WGS84 decimal degrees.
    epsg: 'EPSG:25830' (ETRS89 UTM 30N), 'EPSG:32630' (WGS84 UTM 30N), etc.
    Returns (lat_deg, lon_deg) or None if EPSG not recognized.
    """
    import math

    zone = None
    is_north = True
    code = str(epsg).upper().replace("EPSG:", "").strip()
    try:
        epsg_num = int(code)
        if 25828 <= epsg_num <= 25838:          # ETRS89 UTM zones 28-38N
            zone = epsg_num - 25800
        elif 32601 <= epsg_num <= 32660:         # WGS84 UTM zones 1-60N
            zone = epsg_num - 32600
        elif 32701 <= epsg_num <= 32760:         # WGS84 UTM zones 1-60S
            zone = epsg_num - 32700
            is_north = False
    except (ValueError, TypeError):
        pass

    if zone is None or not (1 <= zone <= 60):
        return None

    # WGS84 ellipsoid constants
    a   = 6378137.0
    f   = 1 / 298.257223563
    b   = a * (1 - f)
    e2  = 1 - (b / a) ** 2
    ep2 = (a / b) ** 2 - 1
    k0  = 0.9996
    E0  = 500000.0
    N0  = 0.0 if is_north else 10000000.0
    lon0 = math.radians((zone - 1) * 6 - 180 + 3)

    x = easting - E0
    y = northing - N0
    M = y / k0
    mu = M / (a * (1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256))

    e1 = (1 - math.sqrt(1 - e2)) / (1 + math.sqrt(1 - e2))
    phi1 = (mu
            + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * math.sin(2 * mu)
            + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * math.sin(4 * mu)
            + (151 * e1 ** 3 / 96) * math.sin(6 * mu)
            + (1097 * e1 ** 4 / 512) * math.sin(8 * mu))

    sin_p = math.sin(phi1); cos_p = math.cos(phi1); tan_p = math.tan(phi1)
    N1 = a / math.sqrt(1 - e2 * sin_p ** 2)
    T1 = tan_p ** 2
    C1 = ep2 * cos_p ** 2
    R1 = a * (1 - e2) / (1 - e2 * sin_p ** 2) ** 1.5
    D  = x / (N1 * k0)

    lat = phi1 - (N1 * tan_p / R1) * (
        D ** 2 / 2
        - (5 + 3 * T1 + 10 * C1 - 4 * C1 ** 2 - 9 * ep2) * D ** 4 / 24
        + (61 + 90 * T1 + 298 * C1 + 45 * T1 ** 2 - 252 * ep2 - 3 * C1 ** 2) * D ** 6 / 720)
    lon = lon0 + (
        D
        - (1 + 2 * T1 + C1) * D ** 3 / 6
        + (5 - 2 * C1 + 28 * T1 - 3 * C1 ** 2 + 8 * ep2 + 24 * T1 ** 2) * D ** 5 / 120
    ) / cos_p

    return math.degrees(lat), math.degrees(lon)


def _get_site_coords(ifc) -> dict:
    """Extrae RefLatitude, RefLongitude, RefElevation del IfcSite.
    Fallback: si IfcSite no tiene RefLatitude (Civil 3D con CRS proyectado),
    convierte el punto de emplazamiento UTM → WGS84.
    """
    try:
        sites = ifc.by_type("IfcSite")
        if not sites:
            return {}
        site = sites[0]

        def _dms_to_decimal(dms):
            if dms is None:
                return None
            if isinstance(dms, (int, float)):
                return float(dms)
            try:
                parts = list(dms)
                d  = float(parts[0]) if parts else 0
                m  = float(parts[1]) / 60          if len(parts) > 1 else 0
                s  = float(parts[2]) / 3600         if len(parts) > 2 else 0
                ms = float(parts[3]) / 3600000000   if len(parts) > 3 else 0
                return d + m + s + ms
            except Exception:
                return None

        result = {
            "RefLatitude":  _dms_to_decimal(getattr(site, "RefLatitude", None)),
            "RefLongitude": _dms_to_decimal(getattr(site, "RefLongitude", None)),
            "RefElevation": getattr(site, "RefElevation", None),
            "coords_x": None,
            "coords_y": None,
            "coords_z": None,
        }

        # Coordenadas locales crudas del emplazamiento (ObjectPlacement del IfcSite),
        # antes de cualquier conversión — se exponen tal cual como "Traslación.X/Y/Z".
        coords_list: list | None = None
        try:
            placement = getattr(site, "ObjectPlacement", None)
            if placement and placement.is_a("IfcLocalPlacement"):
                ax = getattr(placement, "RelativePlacement", None)
                if ax and ax.is_a("IfcAxis2Placement3D"):
                    pt = getattr(ax, "Location", None)
                    if pt and pt.is_a("IfcCartesianPoint"):
                        coords_list = list(pt.Coordinates)
        except Exception:
            coords_list = None

        if coords_list:
            if len(coords_list) > 0:
                result["coords_x"] = float(coords_list[0])
            if len(coords_list) > 1:
                result["coords_y"] = float(coords_list[1])
            if len(coords_list) > 2:
                result["coords_z"] = float(coords_list[2])

        # Fallback para Civil 3D: leer la posición UTM del IfcSite y convertir a WGS84
        if (result["RefLatitude"] is None or result["RefLongitude"] is None) and coords_list and len(coords_list) >= 2:
            try:
                crses = ifc.by_type("IfcProjectedCRS")
                if crses:
                    epsg = getattr(crses[0], "Name", "") or ""
                    wgs = _utm_to_wgs84(
                        float(coords_list[0]),
                        float(coords_list[1]),
                        epsg,
                    )
                    if wgs:
                        lat_d, lon_d = wgs
                        if result["RefLatitude"] is None:
                            result["RefLatitude"] = lat_d
                        if result["RefLongitude"] is None:
                            result["RefLongitude"] = lon_d
                        if result["RefElevation"] is None and len(coords_list) > 2:
                            result["RefElevation"] = float(coords_list[2])
            except Exception:
                pass

        return result
    except Exception:
        return {}


def _get_psets_for_entity(entity, import_eu=True) -> dict:
    """Obtiene todos los psets de una entidad con get_psets()."""
    try:
        import ifcopenshell.util.element as eu
        return eu.get_psets(entity)
    except Exception:
        return {}


def _material_name_from_usage(mat_usage) -> str:
    if mat_usage is None:
        return ""
    # IfcMaterial directo
    if mat_usage.is_a("IfcMaterial"):
        return mat_usage.Name or ""
    # IfcMaterialLayerSetUsage → IfcMaterialLayerSet
    if mat_usage.is_a("IfcMaterialLayerSetUsage"):
        ls = getattr(mat_usage, "ForLayerSet", None)
        if ls:
            layers = getattr(ls, "MaterialLayers", []) or []
            names = [
                getattr(getattr(l, "Material", None), "Name", "") or ""
                for l in layers
            ]
            return "; ".join(n for n in names if n)
    # IfcMaterialLayerSet
    if mat_usage.is_a("IfcMaterialLayerSet"):
        layers = getattr(mat_usage, "MaterialLayers", []) or []
        names = [
            getattr(getattr(l, "Material", None), "Name", "") or ""
            for l in layers
        ]
        return "; ".join(n for n in names if n)
    # IfcMaterialList
    if mat_usage.is_a("IfcMaterialList"):
        mats = getattr(mat_usage, "Materials", []) or []
        return "; ".join(getattr(m, "Name", "") or "" for m in mats)
    return ""


def _classification_from_ref(ref) -> dict:
    if ref is None:
        return {"system": "", "reference": "", "name": ""}
    system = ""
    try:
        src = getattr(ref, "ReferencedSource", None)
        if src:
            system = getattr(src, "Name", "") or ""
    except Exception:
        pass
    return {
        "system":    system,
        "reference": getattr(ref, "Identification", None) or getattr(ref, "ItemReference", None) or "",
        "name":      getattr(ref, "Name", "") or "",
    }


def _build_entity_indexes(ifc) -> tuple[dict, dict, dict]:
    """
    Construye, en UNA sola pasada por cada tipo de relación, los mapas
    entity.id() → material / clasificación / contenedor espacial.

    Antes esto se resolvía dentro del bucle de elementos volviendo a
    recorrer TODAS las relaciones del archivo por cada elemento
    (O(elementos × relaciones)) — con archivos grandes eso hacía que el
    escaneo pareciera colgado. Precomputar aquí lo deja en O(elementos + relaciones).
    """
    material_by_id: dict = {}
    try:
        for rel in ifc.by_type("IfcRelAssociatesMaterial"):
            name = _material_name_from_usage(getattr(rel, "RelatingMaterial", None))
            if not name:
                continue
            for obj in (getattr(rel, "RelatedObjects", []) or []):
                material_by_id[obj.id()] = name
    except Exception:
        pass

    classification_by_id: dict = {}
    try:
        for rel in ifc.by_type("IfcRelAssociatesClassification"):
            cls = _classification_from_ref(getattr(rel, "RelatingClassification", None))
            for obj in (getattr(rel, "RelatedObjects", []) or []):
                classification_by_id[obj.id()] = cls
    except Exception:
        pass

    spatial_by_id: dict = {}
    try:
        for rel in ifc.by_type("IfcRelContainedInSpatialStructure"):
            container = getattr(rel, "RelatingStructure", None)
            cname = getattr(container, "Name", "") or "" if container else ""
            for obj in (getattr(rel, "RelatedElements", []) or []):
                spatial_by_id[obj.id()] = cname
    except Exception:
        pass

    return material_by_id, classification_by_id, spatial_by_id


def _extract_cell(source: str, entity, psets: dict, site_coords: dict,
                  material: str, classification: dict, file_info: dict) -> str:
    """Extrae el valor de una celda según la definición de source."""
    if source == "material":
        return material

    if source.startswith("classification."):
        field = source.split(".", 1)[1]
        return classification.get(field, "")

    if source.startswith("info."):
        field = source.split(".", 1)[1]
        return str(file_info.get(field, "") or "")

    if source.startswith("attr."):
        attr = source.split(".", 1)[1]
        if attr == "ifc_class":
            return entity.is_a() if entity else ""
        if attr == "spatial_container":
            return file_info.get("_spatial_container_" + str(getattr(entity, "GlobalId", "")), "")
        value = str(getattr(entity, attr, "") or "") if entity else ""
        if attr == "LongName":
            # IfcProject.LongName en Revit sale del campo nativo "Nombre de
            # proyecto", que varía sin control de un .rvt a otro (y el
            # override por parámetro personalizado no se respeta en
            # Revit 2025+). Usamos siempre el nombre de actuación ADIF en su
            # lugar, que sí llega consistente desde Revit y desde Civil 3D.
            fallback = psets.get("ADIF_02_Localizacion", {}).get("ADIF_02_Entorno_Actuacion")
            if fallback:
                return str(fallback)
        return value

    if source.startswith("site."):
        field = source.split(".", 1)[1]
        val = site_coords.get(field)
        if val is None:
            return ""
        return str(round(val, 6)) if isinstance(val, float) else str(val)

    if source.startswith("pset:"):
        rest = source[5:]
        pset_name, prop_name = rest.split("/", 1)
        pset_data = psets.get(pset_name, {})
        val = pset_data.get(prop_name)
        if val is None:
            return ""
        return str(val)

    return ""


def _make_cell(value: str, rule_fn) -> dict:
    red, reason = rule_fn(value)
    cell = {"v": value, "red": red}
    if red and reason:
        cell["reason"] = reason
    return cell


# ── Scan en vivo (un archivo) ─────────────────────────────────────────────────

def scan_file(path: str, levels: list[str] | None = None) -> dict:
    """
    Escanea un IFC y devuelve la matriz de revisión del cliente.
    Un request = un archivo. El frontend itera todos los archivos seleccionados.
    """
    if levels is None:
        levels = ["file", "project", "element"]

    cached = _cache_read(path)
    if cached is not None:
        return cached

    if not os.path.isfile(path):
        return {"error": f"File not found: {path}", "path": path}

    try:
        import ifcopenshell
        import ifcopenshell.util.element as eu
    except ImportError:
        return {"error": "ifcopenshell not installed", "path": path}

    try:
        ifc = ifcopenshell.open(path)
    except Exception as e:
        return {"error": str(e), "path": path}

    filename = os.path.basename(path)
    stat = os.stat(path)

    # Info básica del archivo (sin abrir IFC completo de nuevo para algunos campos)
    file_info = {
        "filename": filename,
        "path": path,
        "file_size_bytes": stat.st_size,
        "unidad": "",
        "schema": "",
        "timestamp": "",
        "originating_system": "",
        "app_version": "",
        "organization": "",
    }

    # Extraer header del IFC con el parser STEP probado (maneja Civil 3D multilínea)
    if _extract_ifc_info:
        try:
            ifc_info = _extract_ifc_info(path)
            file_info["schema"]              = ifc_info.get("schema", "")
            file_info["organization"]        = ifc_info.get("organization", "")
            file_info["originating_system"]  = ifc_info.get("originating_system", "")
            file_info["app_version"]         = ifc_info.get("app_version", "")
            file_info["timestamp"]           = ifc_info.get("timestamp", "")
        except Exception as e:
            print(f"[client_review] {filename}: ERROR extracting ifc_info: {e}", file=sys.stderr)

    site_coords = _get_site_coords(ifc)

    result: dict = {"path": path, "filename": filename}

    # ── Nivel archivo ─────────────────────────────────────────────────────────
    if "file" in levels:
        file_row: dict[str, dict] = {}
        for col in SCHEMA["file"]:
            val = _extract_cell(col["source"], None, {}, site_coords, "", {}, file_info)
            file_row[col["key"]] = _make_cell(val, col["rule"])
        result["file_row"] = file_row

    # ── Nivel proyecto ────────────────────────────────────────────────────────
    if "project" in levels:
        project_row: dict[str, dict] = {}
        project_guid = ""
        try:
            projects = ifc.by_type("IfcProject")
            if projects:
                proj = projects[0]
                project_guid = getattr(proj, "GlobalId", "") or ""
                psets = _get_psets_for_entity(proj)
                for col in SCHEMA["project"]:
                    val = _extract_cell(col["source"], proj, psets, site_coords, "", {}, file_info)
                    project_row[col["key"]] = _make_cell(val, col["rule"])
        except Exception as e:
            print(f"[client_review] project scan error {filename}: {e}", file=sys.stderr)
        result["project_row"] = project_row
        result["project_guid"] = project_guid

    # ── Nivel elemento ────────────────────────────────────────────────────────
    if "element" in levels:
        elements = []
        red_by_col: dict[str, int] = defaultdict(int)
        by_disciplina: dict[str, int] = defaultdict(int)
        by_ifc_class: dict[str, int] = defaultdict(int)
        red_total = 0

        try:
            all_entities = list(ifc.by_type("IfcProduct"))
            material_by_id, classification_by_id, spatial_by_id = _build_entity_indexes(ifc)
            empty_classification = {"system": "", "reference": "", "name": ""}

            for entity in all_entities:
                guid = getattr(entity, "GlobalId", "") or ""
                psets = _get_psets_for_entity(entity)
                eid = entity.id()
                material = material_by_id.get(eid, "")
                classification = classification_by_id.get(eid, empty_classification)
                file_info["_spatial_container_" + guid] = spatial_by_id.get(eid, "")

                cells: dict[str, dict] = {}
                elem_red = 0
                for col in SCHEMA["element"]:
                    val = _extract_cell(col["source"], entity, psets, site_coords,
                                        material, classification, file_info)
                    cell = _make_cell(val, col["rule"])
                    cells[col["key"]] = cell
                    if cell["red"]:
                        red_by_col[col["key"]] += 1
                        elem_red += 1
                        red_total += 1

                disciplina = ""
                adif00 = psets.get("ADIF_00_Generales", {})
                if adif00:
                    disciplina = str(adif00.get("ADIF_00_Disciplina", "") or "")
                if disciplina:
                    by_disciplina[disciplina] += elem_red
                by_ifc_class[entity.is_a()] += 1

                elements.append({
                    "guid":     guid,
                    "filename": filename,
                    "cells":    cells,
                })
        except Exception as e:
            print(f"[client_review] element scan error {filename}: {e}", file=sys.stderr)

        result["elements"] = elements
        result["stats"] = {
            "element_count": len(elements),
            "red_total":     red_total,
            "red_by_col":    dict(red_by_col),
            "by_disciplina": dict(by_disciplina),
            "by_ifc_class":  dict(by_ifc_class),
        }

    _cache_write(path, result)
    return result


# ── Importar Excel del cliente ────────────────────────────────────────────────

def import_excel(file_bytes: bytes, filename: str, review_id: str) -> dict:
    """
    Parsea el xlsx del cliente detectando celdas rojas (fgColor.rgb == FFFF0000).
    Acepta también CSV (sin información de color).
    Persiste el resultado a disco bajo _REVIEWS_DIR/<review_id>.json.
    """
    ext = os.path.splitext(filename)[1].lower()
    is_csv = ext == ".csv"

    sheets_data: dict[str, dict] = {}

    if is_csv:
        sheets_data = _parse_csv(file_bytes, filename)
    else:
        sheets_data = _parse_xlsx(file_bytes)

    # Extraer lista de archivos IFC mencionados en las hojas
    files_in_excel: list[str] = []
    for sheet_name, sheet in sheets_data.items():
        for row in sheet.get("rows", []):
            # Columna "Nombre" en IfcFile, "Archivo" en IfcProject/IfcElement
            for col_key in ("Nombre", "Archivo", "filename"):
                val = row.get(col_key, {})
                if isinstance(val, dict):
                    val = val.get("v", "")
                if val and val not in files_in_excel:
                    files_in_excel.append(val)

    result = {
        "review_id": review_id,
        "source_filename": filename,
        "sheets": {
            name: {
                "rows": len(sheet.get("rows", [])),
                "red_cells": sum(
                    1 for row in sheet.get("rows", [])
                    for cell in row.values()
                    if isinstance(cell, dict) and cell.get("red")
                ),
                "columns": sheet.get("columns", []),
            }
            for name, sheet in sheets_data.items()
        },
        "files_in_excel": files_in_excel,
        "_data": sheets_data,  # datos completos, no se devuelven al frontend en /import
    }

    # Persistir a disco
    os.makedirs(_REVIEWS_DIR, exist_ok=True)
    fp = os.path.join(_REVIEWS_DIR, f"{review_id}.json")
    tmp = fp + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(result, fh, ensure_ascii=False)
    os.replace(tmp, fp)

    return {k: v for k, v in result.items() if k != "_data"}


def _normalize_header(raw: str) -> str:
    """Normaliza cabeceras de Excel: elimina _x000D_, divide por \\n y toma el último segmento."""
    s = raw.replace("_x000D_", "").replace("\r", "")
    parts = [p.strip() for p in s.split("\n") if p.strip()]
    return parts[-1] if parts else s.strip()


def _parse_xlsx(file_bytes: bytes) -> dict:
    """Parsea xlsx y marca celdas rojas."""
    import io
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets: dict[str, dict] = {}

    for ws in wb.worksheets:
        sname = ws.title
        rows_data = list(ws.iter_rows(values_only=False))
        if not rows_data:
            sheets[sname] = {"columns": [], "rows": []}
            continue

        # Detectar fila de cabecera (primera fila no vacía)
        header_idx = 0
        columns: list[str] = []
        for i, row in enumerate(rows_data):
            vals = [c.value for c in row if c.value is not None]
            if vals:
                columns = [_normalize_header(str(c.value or "")) for c in row]
                header_idx = i
                break

        parsed_rows: list[dict] = []
        for row in rows_data[header_idx + 1:]:
            row_dict: dict[str, dict] = {}
            for col_idx, cell in enumerate(row):
                if col_idx >= len(columns):
                    break
                col_name = columns[col_idx]
                if not col_name:
                    continue
                val = str(cell.value or "") if cell.value is not None else ""
                is_red = False
                try:
                    fill = cell.fill
                    if (fill and fill.patternType == "solid"
                            and fill.fgColor
                            and str(fill.fgColor.rgb).upper() in ("FFFF0000", "FF0000")):
                        is_red = True
                except Exception:
                    pass
                row_dict[col_name] = {"v": val, "red": is_red}
            if any(c.get("v") or c.get("red") for c in row_dict.values()):
                parsed_rows.append(row_dict)

        sheets[sname] = {"columns": [c for c in columns if c], "rows": parsed_rows}

    return sheets


def _parse_csv(file_bytes: bytes, filename: str) -> dict:
    """Parsea un CSV individual (sin info de color, solo vacíos = red)."""
    import csv
    import io

    sheet_name = os.path.splitext(filename)[0]
    text = file_bytes.decode("utf-8", errors="replace")

    # Detectar separador
    sep = ";" if text.count(";") > text.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=sep)
    raw_rows = list(reader)
    if not raw_rows:
        return {sheet_name: {"columns": [], "rows": []}}

    columns = [_normalize_header(c) for c in raw_rows[0]]
    parsed_rows: list[dict] = []
    for raw in raw_rows[1:]:
        row_dict: dict[str, dict] = {}
        for i, col in enumerate(columns):
            if not col:
                continue
            val = raw[i].strip() if i < len(raw) else ""
            is_red = val.lower() in EMPTY_VALUES or val == ""
            row_dict[col] = {"v": val, "red": is_red}
        if any(c.get("v") for c in row_dict.values()):
            parsed_rows.append(row_dict)

    return {sheet_name: {"columns": [c for c in columns if c], "rows": parsed_rows}}


def _load_review(review_id: str) -> dict | None:
    """Carga un review persistido por import_excel."""
    fp = os.path.join(_REVIEWS_DIR, f"{review_id}.json")
    try:
        with open(fp, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, json.JSONDecodeError):
        return None


# ── Comparar Excel del cliente vs IFC en disco ────────────────────────────────

def compare_file(path: str, review_id: str) -> dict:
    """
    Compara el escaneo en vivo de un IFC contra los datos importados del cliente.
    Devuelve misma estructura que scan_file pero con status por celda:
      client_red_still_empty  — rojo cliente y sigue vacío
      client_red_fixed        — rojo cliente, ahora tiene valor
      client_red_has_value    — rojo cliente pero ya tenía valor (discrepancia)
      ok                      — sin rojo y con valor
      new_empty               — vacío en disco no marcado por cliente (posible regresión)
      not_in_excel            — elemento no estaba en la revisión del cliente
    """
    review = _load_review(review_id)
    if review is None:
        return {"error": f"Review {review_id} not found", "path": path}

    scan = scan_file(path)
    if "error" in scan:
        return scan

    filename = os.path.basename(path)
    sheets = review.get("_data", {})

    def _compare_cell(disk_cell: dict, client_cell: dict | None) -> dict:
        disk_val = disk_cell.get("v", "")
        disk_empty = _is_empty(disk_val)

        if client_cell is None:
            cell = dict(disk_cell)
            cell["status"] = "new_empty" if disk_empty else "ok"
            return cell

        client_red = client_cell.get("red", False)
        client_val = client_cell.get("v", "")

        cell = dict(disk_cell)
        if client_red:
            if disk_empty:
                cell["status"] = "client_red_still_empty"
                cell["client_value"] = client_val
            elif not _is_empty(client_val) and client_val != disk_val:
                cell["status"] = "client_red_has_value"
                cell["client_value"] = client_val
            else:
                cell["status"] = "client_red_fixed"
                cell["client_value"] = client_val
        else:
            cell["status"] = "new_empty" if disk_empty else "ok"
        return cell

    # ── Nivel archivo ─────────────────────────────────────────────────────────
    file_row = scan.get("file_row", {})
    client_file_row: dict = {}
    # Buscar en IfcFile sheet por filename (case-insensitive)
    for sheet_name in ("IfcFile", "ifcFile", "ifcfile"):
        sheet = sheets.get(sheet_name, {})
        for row in sheet.get("rows", []):
            row_fname = (row.get("Nombre", {}) or {}).get("v", "")
            if row_fname.lower() == filename.lower():
                client_file_row = row
                break
        if client_file_row:
            break

    compared_file_row = {
        col: _compare_cell(file_row[col], client_file_row.get(col))
        for col in file_row
    }

    # ── Nivel proyecto ────────────────────────────────────────────────────────
    project_row = scan.get("project_row", {})
    project_guid = scan.get("project_guid", "")
    client_project_row: dict = {}
    for sheet_name in ("IfcProject", "ifcProject", "ifcproject"):
        sheet = sheets.get(sheet_name, {})
        for row in sheet.get("rows", []):
            row_guid = (row.get("GUID", {}) or {}).get("v", "")
            if row_guid and row_guid == project_guid:
                client_project_row = row
                break
        if client_project_row:
            break

    compared_project_row = {
        col: _compare_cell(project_row[col], client_project_row.get(col))
        for col in project_row
    }

    # ── Nivel elemento ────────────────────────────────────────────────────────
    # Indexar filas del cliente por GUID
    client_elements: dict[str, dict] = {}
    for sheet_name in ("IfcElement", "ifcElement", "ifcelement"):
        sheet = sheets.get(sheet_name, {})
        for row in sheet.get("rows", []):
            row_filename = (row.get("Archivo", {}) or {}).get("v", "")
            if row_filename and row_filename.lower() != filename.lower():
                continue
            row_guid = (row.get("IfcGUID", {}) or {}).get("v", "")
            if row_guid:
                client_elements[row_guid] = row
        if client_elements:
            break

    compared_elements = []
    for elem in scan.get("elements", []):
        guid = elem["guid"]
        client_elem = client_elements.get(guid)
        status_prefix = "not_in_excel" if client_elem is None else ""
        cells = elem["cells"]
        compared_cells = {
            col: _compare_cell(cells[col], (client_elem or {}).get(col))
            for col in cells
        }
        if status_prefix:
            for col in compared_cells:
                compared_cells[col]["status"] = "not_in_excel"
        compared_elements.append({
            "guid":     guid,
            "filename": filename,
            "cells":    compared_cells,
        })

    # Stats
    status_counts: dict[str, int] = defaultdict(int)
    for elem in compared_elements:
        for cell in elem["cells"].values():
            s = cell.get("status", "ok")
            status_counts[s] += 1

    return {
        "path":               path,
        "filename":           filename,
        "file_row":           compared_file_row,
        "project_row":        compared_project_row,
        "project_guid":       project_guid,
        "elements":           compared_elements,
        "client_elements_matched": len(client_elements),
        "stats": {
            **scan.get("stats", {}),
            "status_counts": dict(status_counts),
        },
    }


# ── Export / Import Excel ─────────────────────────────────────────────────────

def export_to_excel(data: dict) -> bytes:
    """
    Crea un xlsx limpio (sin colores, solo negritas en cabeceras) desde los datos de la tabla.
    data: {sheetName: {columns: [...], rows: [[...]]}}
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, sheet_data in data.items():
        ws = wb.create_sheet(title=str(sheet_name)[:31])
        columns = sheet_data.get("columns", [])
        rows    = sheet_data.get("rows", [])

        if columns:
            ws.append(columns)
            # Las cabeceras vienen como "Grupo\nSubcabecera" (igual que el
            # Excel del cliente) — wrap_text para que el \n se vea como
            # salto de línea real, no como caracter suelto.
            has_multiline = any("\n" in str(c) for c in columns)
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if has_multiline:
                ws.row_dimensions[1].height = 30

        for row in rows:
            ws.append([str(v) if v is not None else "" for v in row])

        # Auto-width aproximado (basado en la línea más larga de cada celda)
        for col_cells in ws.columns:
            max_len = max(
                (max((len(ln) for ln in str(c.value or "").split("\n")), default=0) for c in col_cells),
                default=0,
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def apply_excel_edits(file_bytes: bytes, paths_by_filename: dict[str, str]) -> dict:
    """
    Parsea un xlsx exportado por nosotros y escribe los valores editables de vuelta
    a los archivos IFC correspondientes.
    Solo aplica columnas con source "pset:..." en el schema (campos editables).
    Abre cada IFC una sola vez y aplica todos sus cambios.
    """
    import io
    import ifcopenshell
    import ifcopenshell.guid as ifc_guid
    import ifcopenshell.util.element as eu
    from openpyxl import load_workbook
    from services.client_review_schema import SCHEMA

    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)

    # col_key → (pset_name, prop_name) para columnas editables
    def _pset_info(col: str, level: str):
        for c in SCHEMA.get(level, []):
            if c["key"] == col and c["source"].startswith("pset:"):
                rest = c["source"][5:]
                pn, pp = rest.split("/", 1)
                return pn, pp
        return None

    overall: dict[str, dict] = {}

    sheet_map = [
        ("IfcElement", "element", "IfcGUID"),
        ("IfcProject", "project", "GUID"),
    ]

    for sheet_name, level, guid_col in sheet_map:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            continue

        headers = [str(c).strip() if c else "" for c in all_rows[0]]
        if guid_col not in headers:
            continue

        file_col = "Archivo"

        # Agrupar filas por ruta completa
        rows_by_path: dict[str, list[dict]] = {}
        for raw in all_rows[1:]:
            if not any(raw):
                continue
            row_dict = {
                headers[i]: (str(raw[i]) if raw[i] is not None else "")
                for i in range(min(len(headers), len(raw)))
            }
            filename = row_dict.get(file_col, "")
            path = paths_by_filename.get(filename)
            if not path or not os.path.isfile(path):
                continue
            rows_by_path.setdefault(path, []).append(row_dict)

        for path, file_rows in rows_by_path.items():
            filename = os.path.basename(path)
            if filename not in overall:
                overall[filename] = {"updated": 0, "skipped": 0, "errors": []}

            try:
                ifc = ifcopenshell.open(path)
                wrote = False

                for row_dict in file_rows:
                    guid = row_dict.get(guid_col, "")
                    if not guid:
                        continue

                    try:
                        entity = ifc.by_guid(guid)
                    except Exception:
                        entity = None

                    if entity is None:
                        overall[filename]["errors"].append(f"GUID no encontrado: {guid[:12]}")
                        continue

                    current_psets = eu.get_psets(entity)

                    # Pre-indexar psets existentes
                    existing_psets: dict[str, object] = {}
                    for rel in ifc.by_type("IfcRelDefinesByProperties"):
                        try:
                            related = getattr(rel, "RelatedObjects", None) or []
                            if entity not in related:
                                continue
                            pd = rel.RelatingPropertyDefinition
                            if pd and pd.is_a("IfcPropertySet"):
                                existing_psets[pd.Name] = pd
                        except Exception:
                            continue

                    for col, new_val in row_dict.items():
                        pi = _pset_info(col, level)
                        if pi is None:
                            continue
                        pset_name, prop_name = pi

                        current = str((current_psets.get(pset_name) or {}).get(prop_name, "") or "")
                        if current == new_val:
                            overall[filename]["skipped"] += 1
                            continue

                        try:
                            pset = existing_psets.get(pset_name)
                            if pset is None:
                                pset = ifc.create_entity(
                                    "IfcPropertySet",
                                    GlobalId=ifc_guid.new(),
                                    OwnerHistory=None,
                                    Name=pset_name,
                                    HasProperties=(),
                                )
                                ifc.create_entity(
                                    "IfcRelDefinesByProperties",
                                    GlobalId=ifc_guid.new(),
                                    OwnerHistory=None,
                                    RelatedObjects=[entity],
                                    RelatingPropertyDefinition=pset,
                                )
                                existing_psets[pset_name] = pset

                            props = list(getattr(pset, "HasProperties", None) or [])
                            existing_prop = next((p for p in props if p.Name == prop_name), None)

                            if existing_prop is not None:
                                nv = getattr(existing_prop, "NominalValue", None)
                                if nv is not None:
                                    nv.wrappedValue = new_val
                                else:
                                    existing_prop.NominalValue = ifc.create_entity("IfcText", new_val)
                            else:
                                new_p = ifc.create_entity(
                                    "IfcPropertySingleValue",
                                    Name=prop_name,
                                    Description=None,
                                    NominalValue=ifc.create_entity("IfcText", new_val),
                                    Unit=None,
                                )
                                pset.HasProperties = props + [new_p]

                            overall[filename]["updated"] += 1
                            wrote = True
                        except Exception as e2:
                            overall[filename]["errors"].append(f"{col}: {str(e2)[:80]}")

                if wrote:
                    atomic_write_model(ifc, path)

            except Exception as e:
                overall[filename]["errors"].append(str(e)[:120])

    invalidate_cache(list(paths_by_filename.values()))
    return overall


# ── Relleno masivo ────────────────────────────────────────────────────────────

def bulk_fill(targets_by_file: dict[str, list[str]],
              fills: list[dict],
              ifc_class: str | None = None,
              name_base: str | None = None) -> dict:
    """
    Abre cada IFC una vez y aplica N cambios (evita re-abrir por cada celda).
    targets_by_file: {path: [guid,...]} — usar "__all__" para todos los elementos.
    fills: [{pset_name, prop_name, value, only_if_empty, create_if_missing}]
    ifc_class: filtro adicional server-side por clase IFC (e.g. "IfcWall")
    name_base: filtro por prefijo de nombre sin sufijo _\\d+
    """
    import ifcopenshell
    import ifcopenshell.guid

    re_suffix = re.compile(r"_\d+$")

    per_file: dict = {}
    total_updated = 0
    total_created = 0
    total_skipped = 0

    for path, guid_list in targets_by_file.items():
        if not os.path.isfile(path):
            per_file[os.path.basename(path)] = {"error": "File not found"}
            continue
        try:
            ifc = ifcopenshell.open(path)
        except Exception as e:
            per_file[os.path.basename(path)] = {"error": str(e)}
            continue

        all_entities = list(ifc.by_type("IfcProduct"))
        import ifcopenshell.util.element as eu

        use_all = "__all__" in guid_list
        guid_set = set(guid_list) if not use_all else None

        file_updated = 0
        file_created = 0
        file_skipped = 0
        wrote = False

        for entity in all_entities:
            guid = getattr(entity, "GlobalId", "") or ""

            # Filtro por GUIDs
            if not use_all and guid not in guid_set:
                continue

            # Filtro por clase IFC
            if ifc_class and not entity.is_a(ifc_class):
                continue

            # Filtro por nombre base
            if name_base:
                ename = getattr(entity, "Name", "") or ""
                ename_base = re_suffix.sub("", ename)
                if ename_base != name_base:
                    continue

            psets = eu.get_psets(entity)

            # Pre-indexar psets existentes para este elemento
            existing_psets: dict[str, object] = {}
            for rel in ifc.by_type("IfcRelDefinesByProperties"):
                try:
                    related = getattr(rel, "RelatedObjects", None) or []
                    if entity not in related:
                        continue
                    pd = rel.RelatingPropertyDefinition
                    if pd and pd.is_a("IfcPropertySet"):
                        existing_psets[pd.Name] = pd
                except Exception:
                    continue

            for fill in fills:
                pset_name       = fill.get("pset_name", "")
                prop_name       = fill.get("prop_name", "")
                value           = fill.get("value", "")
                only_if_empty   = fill.get("only_if_empty", True)
                create_if_miss  = fill.get("create_if_missing", True)

                if not pset_name or not prop_name:
                    continue

                # Atributo estándar
                if pset_name == "__attr__":
                    editable = {"Name", "Description", "ObjectType", "Tag"}
                    if prop_name not in editable:
                        continue
                    current = str(getattr(entity, prop_name, "") or "")
                    if only_if_empty and not _is_empty(current):
                        file_skipped += 1
                        continue
                    setattr(entity, prop_name, value or None)
                    file_updated += 1
                    wrote = True
                    continue

                # PropertySet
                pset = existing_psets.get(pset_name)

                if pset is None:
                    if not create_if_miss:
                        file_skipped += 1
                        continue
                    pset = ifc.create_entity(
                        "IfcPropertySet",
                        GlobalId=ifcopenshell.guid.new(),
                        OwnerHistory=None,
                        Name=pset_name,
                        HasProperties=(),
                    )
                    ifc.create_entity(
                        "IfcRelDefinesByProperties",
                        GlobalId=ifcopenshell.guid.new(),
                        OwnerHistory=None,
                        RelatedObjects=[entity],
                        RelatingPropertyDefinition=pset,
                    )
                    existing_psets[pset_name] = pset

                current_props = list(getattr(pset, "HasProperties", None) or [])
                existing_prop = next((p for p in current_props if p.Name == prop_name), None)

                if existing_prop is not None:
                    nv = getattr(existing_prop, "NominalValue", None)
                    current_val = ""
                    if nv is not None:
                        current_val = str(getattr(nv, "wrappedValue", "") or "")
                    if only_if_empty and not _is_empty(current_val):
                        file_skipped += 1
                        continue
                    try:
                        if nv is not None:
                            nv.wrappedValue = value
                        else:
                            existing_prop.NominalValue = ifc.create_entity("IfcText", value)
                        file_updated += 1
                        wrote = True
                    except Exception:
                        try:
                            ifc.remove(existing_prop)
                            current_props = [p for p in current_props if p.Name != prop_name]
                            new_p = ifc.create_entity(
                                "IfcPropertySingleValue",
                                Name=prop_name, Description=None,
                                NominalValue=ifc.create_entity("IfcText", value), Unit=None,
                            )
                            pset.HasProperties = current_props + [new_p]
                            file_updated += 1
                            wrote = True
                        except Exception as e2:
                            print(f"[bulk_fill] {prop_name}: {e2}", file=sys.stderr)
                else:
                    if not create_if_miss:
                        file_skipped += 1
                        continue
                    new_p = ifc.create_entity(
                        "IfcPropertySingleValue",
                        Name=prop_name, Description=None,
                        NominalValue=ifc.create_entity("IfcText", value), Unit=None,
                    )
                    pset.HasProperties = current_props + [new_p]
                    file_created += 1
                    wrote = True

        if wrote:
            atomic_write_model(ifc, path)

        per_file[os.path.basename(path)] = {
            "updated": file_updated,
            "created": file_created,
            "skipped": file_skipped,
        }
        total_updated += file_updated
        total_created += file_created
        total_skipped += file_skipped

    # Invalidar caché
    invalidate_cache(list(targets_by_file.keys()))

    return {
        "ok":       True,
        "updated":  total_updated,
        "created":  total_created,
        "skipped":  total_skipped,
        "per_file": per_file,
    }
