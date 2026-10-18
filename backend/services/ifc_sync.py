import os
import ifcopenshell
import ifcopenshell.api

from core.config import settings
from core.logging_config import get_logger
from core.model_pool import model_pool
from core.resilience import atomic_write_model, coerce_property_value, pb_client

logger = get_logger(__name__)

PB_URL = settings.POCKETBASE_URL


def _fetch_all_pb(collection: str, fields: str = "", extra_params: dict | None = None) -> list[dict]:
    """Pagina todos los registros de una colección PocketBase (máx 500 por página).

    Delegado en `PBClient` (timeout, reintento y circuit-breaker): si PB está
    caído devuelve [] o el resultado parcial acumulado, nunca lanza.
    """
    params: dict = {}
    if fields:
        params["fields"] = fields
    if extra_params:
        params.update(extra_params)
    return pb_client.get_records(collection, params=params, fallback=[])


# ── Cache de modelos IFC abiertos ─────────────────────────────────────────────
# Delegado en core.model_pool.IFCModelPool (LRU acotado, thread-safe, TTL,
# invalidación por mtime). Evita OOM con múltiples IFC de 500MB y corrupción
# por aperturas concurrentes.
#
# `_ifc_model_cache` y `_ifc_cache_mtime` se conservan SOLO por compatibilidad:
# services/model_checker.py (fuera de alcance) los importa y hace .pop() tras
# escribir — inocuo, porque el pool invalida por mtime automáticamente.
_ifc_model_cache: dict[str, object] = {}
_ifc_cache_mtime: dict[str, float] = {}   # tracker de mtime para invalidar índices derivados

# Caches de resultados de material (invalidadas automáticamente por cambio de mtime)
_pset_mat_cache: dict[str, dict[int, str]] = {}      # path → {step_id: mat_name}
_entity_type_rows_cache: dict[str, list] = {}         # path → rows
_elements_cache: dict[str, dict[str, list]] = {}      # path → {types_key: rows}


def _open_cached(ifc_path: str):
    """Devuelve el modelo IFC desde el pool LRU global.

    Si el archivo fue modificado en disco (re-exportado desde Revit/etc.) el
    pool lo recarga automáticamente (invalidación por mtime). Aquí además se
    invalida el índice de psets derivado cuando se detecta cambio de mtime.
    """
    try:
        current_mtime = os.path.getmtime(ifc_path)
    except OSError:
        current_mtime = 0.0

    if _ifc_cache_mtime.get(ifc_path) != current_mtime:
        # Archivo cambiado → invalidar caches derivadas de este módulo
        try:
            _pset_index_cache.pop(ifc_path, None)
        except NameError:
            pass  # _pset_index_cache se define más adelante en el módulo
        _pset_mat_cache.pop(ifc_path, None)
        _entity_type_rows_cache.pop(ifc_path, None)
        _elements_cache.pop(ifc_path, None)
        _ifc_cache_mtime[ifc_path] = current_mtime

    return model_pool.get(ifc_path)


def cache_stats() -> dict:
    """Tamaño (nº de archivos) de las cachés derivadas de este módulo.

    A diferencia de `model_pool` (LRU acotado), estas NO tienen límite ni
    TTL — crecen un entry por cada archivo distinto que el usuario visita
    con cualquier herramienta, y solo se liberan invalidando por mtime o
    con `clear_ifc_cache()`. Es la causa más probable de que la RAM del
    backend "se vaya sumando" al pasar de una herramienta a otra sin que
    nada la libere.
    """
    return {
        "pset_index_cache": len(_pset_index_cache),
        "pset_mat_cache": len(_pset_mat_cache),
        "entity_type_rows_cache": len(_entity_type_rows_cache),
        "elements_cache": len(_elements_cache),
    }


def clear_ifc_cache(ifc_path: str | None = None):
    """Invalida el pool de modelos (y caches derivadas) para una ruta o todas."""
    if ifc_path:
        model_pool.invalidate(ifc_path)
        _ifc_model_cache.pop(ifc_path, None)
        _ifc_cache_mtime.pop(ifc_path, None)
        try:
            _pset_index_cache.pop(ifc_path, None)
        except NameError:
            pass
        _pset_mat_cache.pop(ifc_path, None)
        _entity_type_rows_cache.pop(ifc_path, None)
        _elements_cache.pop(ifc_path, None)
    else:
        model_pool.clear()
        _ifc_model_cache.clear()
        _ifc_cache_mtime.clear()
        try:
            _pset_index_cache.clear()
        except NameError:
            pass
        _pset_mat_cache.clear()
        _entity_type_rows_cache.clear()
        _elements_cache.clear()


# ── Caché en disco de tipos/elementos de material (persiste entre reinicios) ─
# El pool LRU en memoria (model_pool) desaparece al reiniciar el backend o al
# expulsar por LRU/TTL: un IFC de 100+ MB paga el parseo COMPLETO de nuevo la
# próxima vez que se navega a la herramienta, aunque nada haya cambiado en
# disco. Estos resultados (tipos de entidad y elementos con material) son
# pequeños (JSON) y no dependen de la geometría, así que se persisten en
# ~/.ifc-hub-v2/material_cache/ igual que la caché de geometría del visor
# (services/viewer_geometry.py) — misma clave ruta+mtime+tamaño, invalidación
# automática si el archivo cambia.
import json as _json
import hashlib as _hashlib_mat

_MATERIAL_CACHE_DIR = os.environ.get(
    "IFC_HUB_MATERIAL_CACHE",
    os.path.join(os.path.expanduser("~"), ".ifc-hub-v2", "material_cache"),
)

# La clave de caché SOLO invalidaba por (ruta, mtime, tamaño) del IFC — nunca
# por la versión del algoritmo de resolución de material. Consecuencia real
# (no hipotética): al corregir la prioridad Pset-vs-IfcMaterial nativo (ver
# _read_pset_material) para 442_MTE_TEMPORAL.ifc, el archivo en disco no
# cambió, así que el resultado (calculado con la lógica VIEJA, antes del fix)
# se siguió sirviendo indefinidamente sin importar cuántas veces se
# reiniciara el backend — el bug parecía "no arreglado" aunque el código sí
# lo estuviera. Súbete este número cada vez que cambie qué calcula
# get_elements_for_types/get_materials_for_elements/get_entity_type_rows,
# para invalidar de un guantazo todo lo cacheado con la lógica anterior.
_MATERIAL_CACHE_VERSION = 2


def _material_cache_key(ifc_path: str, kind: str, extra: str = "") -> str | None:
    try:
        st = os.stat(ifc_path)
    except OSError:
        return None
    raw = (f"v{_MATERIAL_CACHE_VERSION}|{os.path.abspath(ifc_path)}|"
           f"{int(st.st_mtime)}|{st.st_size}|{kind}|{extra}")
    return _hashlib_mat.sha1(raw.encode("utf-8")).hexdigest()


def _material_cache_path(key: str) -> str:
    return os.path.join(_MATERIAL_CACHE_DIR, f"{key}.json")


def _material_cache_read(ifc_path: str, kind: str, extra: str = ""):
    """Lee el resultado cacheado en disco, o None si no existe/es inválido.
    La clave incluye mtime+tamaño: un archivo modificado nunca sirve un valor
    obsoleto, simplemente produce un cache-miss (no hace falta invalidar)."""
    key = _material_cache_key(ifc_path, kind, extra)
    if not key:
        return None
    try:
        with open(_material_cache_path(key), "r", encoding="utf-8") as fh:
            return _json.load(fh)
    except (OSError, ValueError):
        return None


def _material_cache_write(ifc_path: str, kind: str, data, extra: str = "") -> None:
    key = _material_cache_key(ifc_path, kind, extra)
    if not key:
        return
    try:
        os.makedirs(_MATERIAL_CACHE_DIR, exist_ok=True)
        fp = _material_cache_path(key)
        tmp = fp + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            _json.dump(data, fh)
        os.replace(tmp, fp)   # escritura atómica
    except OSError:
        pass


def load_ifc_tree(ifc_path: str) -> tuple[list[dict], dict]:
    """Carga la estructura espacial de un IFC como lista plana de nodos.
    Retorna (nodes, summary).
    """
    model = _open_cached(ifc_path)

    # Mapa de agregación: padre → hijos por IfcRelAggregates
    agg_map: dict[int, list] = {}
    for rel in model.by_type("IfcRelAggregates"):
        pid = rel.RelatingObject.id()
        agg_map.setdefault(pid, [])
        for child in rel.RelatedObjects:
            agg_map[pid].append(child)

    # Mapa de contenencia: espacio → elementos por IfcRelContainedInSpatialStructure
    cont_map: dict[int, list] = {}
    for rel in model.by_type("IfcRelContainedInSpatialStructure"):
        pid = rel.RelatingStructure.id()
        cont_map.setdefault(pid, [])
        for elem in rel.RelatedElements:
            cont_map[pid].append(elem)

    SPATIAL_TYPES = {
        "IfcProject", "IfcSite", "IfcBuilding",
        "IfcBuildingStorey", "IfcSpace", "IfcExternalSpatialElement",
    }

    nodes: list[dict] = []

    def walk(element, indent: int, parent_id: int):
        eid   = element.id()
        etype = element.is_a()
        name  = str(getattr(element, "Name", None) or etype)
        children_agg  = agg_map.get(eid, [])
        children_cont = cont_map.get(eid, [])
        has_children  = bool(children_agg or children_cont)

        nodes.append({
            "step_id":     eid,
            "entity_type": etype,
            "name":        name,
            "indent":      indent,
            "indent_px":   f"{indent * 20}px",
            "visible":     indent == 0,
            "expanded":    False,
            "has_children": has_children,
            "is_spatial":  etype in SPATIAL_TYPES,
            "parent_id":   parent_id,
        })

        for child in children_agg:
            walk(child, indent + 1, eid)
        for elem in children_cont:
            walk(elem, indent + 1, eid)

    for proj in model.by_type("IfcProject"):
        walk(proj, 0, -1)

    summary = {
        "schema":    model.schema,
        "products":  str(len(model.by_type("IfcProduct"))),
        "psets":     str(len(model.by_type("IfcPropertySet"))),
        "nodes":     str(len(nodes)),
    }
    return nodes, summary


def get_element_props(ifc_path: str, step_id: int) -> dict:
    """Lee todas las propiedades de un elemento IFC para las 4 pestañas.
    Retorna {attrs, psets, location_rows, cls_rows, rel_rows}.
    Todas las *_rows son listas planas {is_header, label, value}.
    """
    model = _open_cached(ifc_path)
    element = model.by_id(step_id)
    if element is None:
        return {"attrs": [], "psets": [], "location_rows": [], "cls_rows": [], "rel_rows": []}

    # ── Atributos directos ────────────────────────────────────────────────
    ATTR_ORDER = ["IfcEntity", "GlobalId", "Name", "Description",
                  "ObjectType", "Tag", "PredefinedType"]
    raw_attrs: dict[str, str] = {"IfcEntity": element.is_a()}
    for attr in ("GlobalId", "Name", "Description", "ObjectType", "Tag", "PredefinedType"):
        try:
            v = getattr(element, attr, None)
            if v is not None:
                raw_attrs[attr] = str(v)
        except Exception:
            pass
    attrs = sorted(
        [{"name": k, "value": v} for k, v in raw_attrs.items()],
        key=lambda x: ATTR_ORDER.index(x["name"]) if x["name"] in ATTR_ORDER else 99,
    )

    # ── Psets ─────────────────────────────────────────────────────────────
    psets: list[dict] = []
    if hasattr(element, "IsDefinedBy"):
        for rel in element.IsDefinedBy:
            if not rel.is_a("IfcRelDefinesByProperties"):
                continue
            pd = rel.RelatingPropertyDefinition
            if not pd.is_a("IfcPropertySet"):
                continue
            props = []
            for prop in pd.HasProperties:
                try:
                    val = prop.NominalValue.wrappedValue
                    sval = str(val) if val is not None else "—"
                except Exception:
                    sval = "—"
                props.append({
                    "name":      prop.Name,
                    "value":     sval,
                    "pset_name": pd.Name,
                    "row_id":    f"{pd.Name}#{prop.Name}",
                })
            props.sort(key=lambda p: p["name"])
            psets.append({"group": pd.Name, "props": props})
    psets.sort(key=lambda p: p["group"])

    # ── Localización ──────────────────────────────────────────────────────
    loc_rows: list[dict] = []

    def _hdr(label):
        loc_rows.append({"is_header": True, "label": label, "value": ""})

    def _row(label, value):
        loc_rows.append({"is_header": False, "label": label, "value": str(value) if value is not None else "—"})

    # Jerarquía espacial (ContainedInStructure / Decomposes)
    _hdr("Contenido en")
    hierarchy = []
    try:
        for rel in getattr(element, "ContainedInStructure", []):
            container = rel.RelatingStructure
            hierarchy.append((container.is_a(), getattr(container, "Name", "") or ""))
    except Exception:
        pass
    try:
        for rel in getattr(element, "Decomposes", []):
            parent = rel.RelatingObject
            hierarchy.append((parent.is_a(), getattr(parent, "Name", "") or ""))
    except Exception:
        pass
    if hierarchy:
        for etype, ename in hierarchy:
            _row(etype, ename)
    else:
        _row("—", "Sin contenedor espacial")

    # Coordenadas de placement
    try:
        import ifcopenshell.util.placement
        if hasattr(element, "ObjectPlacement") and element.ObjectPlacement:
            m = ifcopenshell.util.placement.get_local_placement(element.ObjectPlacement)
            _hdr("Coordenadas (locales)")
            _row("X", f"{m[0][3]:.4f} m")
            _row("Y", f"{m[1][3]:.4f} m")
            _row("Z", f"{m[2][3]:.4f} m")
    except Exception as exc:
        # Degradación: sin coordenadas, pero el panel de propiedades se sirve igual
        logger.debug("get_element_props: placement ilegible en #%s (%s)", step_id, exc)

    # Capa (IfcPresentationLayerAssignment)
    layer_found = False
    try:
        for rel in model.by_type("IfcPresentationLayerAssignment"):
            assigned = list(getattr(rel, "AssignedItems", []))
            # también buscar en representaciones del elemento
            rep = getattr(element, "Representation", None)
            rep_items = []
            if rep:
                for r in rep.Representations:
                    rep_items += list(r.Items)
            if element in assigned or any(i in assigned for i in rep_items):
                if not layer_found:
                    _hdr("Capa")
                    layer_found = True
                _row("Layer", rel.Name)
    except Exception as exc:
        logger.debug("get_element_props: capas ilegibles en #%s (%s)", step_id, exc)

    # ── Clasificación ────────────────────────────────────────────────────
    cls_rows: list[dict] = []

    def _chdr(label):
        cls_rows.append({"is_header": True, "label": label, "value": ""})

    def _crow(label, value):
        cls_rows.append({"is_header": False, "label": label, "value": str(value) if value is not None else "—"})

    try:
        for assoc in getattr(element, "HasAssociations", []):
            if not assoc.is_a("IfcRelAssociatesClassification"):
                continue
            relating = assoc.RelatingClassification

            if relating.is_a("IfcClassificationReference"):
                ref   = relating
                cls   = ref.ReferencedSource
            elif relating.is_a("IfcClassification"):
                ref   = None
                cls   = relating
            else:
                continue

            cls_name = getattr(cls, "Name", "?") if cls else "?"
            _chdr(str(cls_name))

            if cls:
                for field, label in [
                    ("Name",        "Nombre"),
                    ("Source",      "Fuente"),
                    ("Edition",     "Edición"),
                    ("EditionDate", "Fecha edición"),
                    ("Description", "Descripción"),
                    ("Location",    "Localización"),
                    ("Specification", "Localización"),
                ]:
                    try:
                        v = getattr(cls, field, None)
                        if v:
                            _crow(label, str(v))
                    except Exception:
                        pass

            if ref:
                _chdr("Referencia")
                for field, label in [
                    ("Identification", "Identificación"),
                    ("Name",           "Nombre"),
                    ("Location",       "Localización"),
                    ("Description",    "Descripción"),
                ]:
                    try:
                        v = getattr(ref, field, None)
                        if v:
                            _crow(label, str(v))
                    except Exception:
                        pass
    except Exception as exc:
        logger.debug("get_element_props: clasificación ilegible en #%s (%s)", step_id, exc)

    if not cls_rows:
        cls_rows.append({"is_header": False, "label": "Sin clasificación", "value": ""})

    # ── Relaciones ───────────────────────────────────────────────────────
    rel_rows: list[dict] = []

    def _rhdr(label):
        rel_rows.append({"is_header": True, "label": label, "value": ""})

    def _rrow(label, value):
        rel_rows.append({"is_header": False, "label": label, "value": str(value) if value is not None else "—"})

    # Tipo de objeto (IfcRelDefinesByType)
    try:
        for rel in getattr(element, "IsTypedBy", []):
            obj_type = rel.RelatingType
            _rhdr("Tipo de objeto")
            _rrow("IfcEntity", obj_type.is_a())
            _rrow("Nombre",    getattr(obj_type, "Name", "") or "—")
            _rrow("GUID",      getattr(obj_type, "GlobalId", "") or "—")
    except Exception:
        pass

    # Material (IfcRelAssociatesMaterial)
    try:
        for assoc in getattr(element, "HasAssociations", []):
            if not assoc.is_a("IfcRelAssociatesMaterial"):
                continue
            mat = assoc.RelatingMaterial
            _rhdr("Material")
            if mat.is_a("IfcMaterial"):
                _rrow("Nombre", mat.Name)
            elif mat.is_a("IfcMaterialLayerSetUsage"):
                for i, layer in enumerate(mat.ForLayerSet.MaterialLayers, 1):
                    lname = getattr(layer.Material, "Name", "?") if layer.Material else "?"
                    _rrow(f"Capa {i}", f"{lname} ({layer.LayerThickness} m)")
            elif mat.is_a("IfcMaterialList"):
                for m in mat.Materials:
                    _rrow("Material", m.Name)
            else:
                _rrow("Tipo", mat.is_a())
    except Exception:
        pass

    # Psets vinculados (IfcRelDefinesByProperties)
    pset_names = [ps["group"] for ps in psets]
    if pset_names:
        _rhdr("Property Sets vinculados")
        for name in pset_names:
            _rrow("Pset", name)

    # Clasificaciones vinculadas
    try:
        cls_names = []
        for assoc in getattr(element, "HasAssociations", []):
            if assoc.is_a("IfcRelAssociatesClassification"):
                rel_cls = assoc.RelatingClassification
                if rel_cls.is_a("IfcClassificationReference"):
                    src = rel_cls.ReferencedSource
                    cls_names.append(getattr(src, "Name", "?") if src else "?")
                elif rel_cls.is_a("IfcClassification"):
                    cls_names.append(getattr(rel_cls, "Name", "?"))
        if cls_names:
            _rhdr("Clasificaciones")
            for name in cls_names:
                _rrow("Sistema", name)
    except Exception:
        pass

    # Contenido en (espacial)
    try:
        for rel in getattr(element, "ContainedInStructure", []):
            container = rel.RelatingStructure
            _rhdr("Estructura espacial")
            _rrow(container.is_a(), getattr(container, "Name", "") or "—")
    except Exception:
        pass

    # Parte de (Decomposes)
    try:
        for rel in getattr(element, "Decomposes", []):
            parent = rel.RelatingObject
            _rhdr("Parte de")
            _rrow(parent.is_a(), getattr(parent, "Name", "") or "—")
    except Exception:
        pass

    if not rel_rows:
        rel_rows.append({"is_header": False, "label": "Sin relaciones", "value": ""})

    return {
        "attrs":         attrs,
        "psets":         psets,
        "location_rows": loc_rows,
        "cls_rows":      cls_rows,
        "rel_rows":      rel_rows,
    }


def _detect_software(ifc_path: str) -> str:
    try:
        model = ifcopenshell.open(ifc_path)
        apps = model.by_type("IfcApplication")
        if apps:
            name = str(apps[0].ApplicationFullName).lower()
            if "revit" in name:
                return "Revit"
            if "civil 3d" in name or "civil3d" in name:
                return "Civil3D"
    except Exception:
        pass
    return "?"


def _get_native_name(ifc_filename: str) -> str | None:
    """Busca el native_name de un IFC en midp_mapping, para la etapa activa del
    proyecto. None si no hay match o PB cae."""
    try:
        items = pb_client.get_records(
            "midp_mapping",
            params={
                "filter": f"(ifc_name='{ifc_filename}'&&etapa='{settings.MIDP_ETAPA_ACTIVA}')",
                "perPage": 1,
            },
            fallback=[],
        )
        return items[0]["native_name"] if items else None
    except Exception as exc:
        logger.debug("_get_native_name(%s): %s", ifc_filename, exc)
        return None


def scan_folder(folder_path: str) -> list[dict]:
    """Escanea una carpeta buscando IFCs. Verifica midp_mapping para cada uno."""
    results = []
    for root, dirs, files in os.walk(folder_path):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fname in sorted(files):
            if not fname.lower().endswith(".ifc"):
                continue
            full_path = os.path.join(root, fname)
            native = _get_native_name(fname)
            software = _detect_software(full_path)
            results.append({
                "path":        full_path,
                "filename":    fname,
                "software":    software,
                "native_name": native or "",
                "status":      "READY" if native else "NO-MAP",
                "selected":    native is not None,
            })
    results.sort(key=lambda x: (x["software"], x["filename"]))
    return results


def get_all_midp_mapping(etapa: str | None = None) -> list[dict]:
    """Lee los registros de midp_mapping de PocketBase (paginado).
    Por defecto filtra por la etapa activa del proyecto (midp_mapping conserva el
    histórico de todas las etapas sin borrar filas). Pasar etapa=None para leer
    la etapa activa configurada, o una cadena vacía "" para leer todas."""
    active = settings.MIDP_ETAPA_ACTIVA if etapa is None else etapa
    extra = {"filter": f"(etapa='{active}')"} if active else None
    return _fetch_all_pb("midp_mapping", extra_params=extra)


def validate_folder(folder_path: str) -> dict:
    """Compara archivos IFC de una carpeta contra midp_mapping.
    Devuelve: matched, unmatched (en carpeta, sin mapping), orphaned (en PB, sin IFC)
    """
    ifc_files = set()
    for root, dirs, files in os.walk(folder_path):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fname in files:
            if fname.lower().endswith(".ifc"):
                ifc_files.add(fname)

    mapping = get_all_midp_mapping()
    mapped_names = {m["ifc_name"]: m for m in mapping}

    matched   = []
    unmatched = []
    for fname in sorted(ifc_files):
        if fname in mapped_names:
            matched.append({
                "ifc_name":    fname,
                "native_name": mapped_names[fname].get("native_name", ""),
                "status":      "MATCHED",
            })
        else:
            unmatched.append({
                "ifc_name":    fname,
                "native_name": "",
                "status":      "NO-MAP",
            })

    folder_set = ifc_files
    orphaned = [
        {"ifc_name": m["ifc_name"], "native_name": m.get("native_name", ""), "status": "ORPHAN"}
        for m in mapping
        if m["ifc_name"] not in folder_set
    ]

    return {"matched": matched, "unmatched": unmatched, "orphaned": orphaned}


def get_pb_params() -> list[dict]:
    """Lee ifc_project_defaults de PocketBase."""
    return _fetch_all_pb("ifc_project_defaults")


def scan_folder_simple(folder_path: str) -> list[dict]:
    """Escanea una carpeta y devuelve todos los IFCs sin validar midp_mapping."""
    results = []
    for root, dirs, files in os.walk(folder_path):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fname in sorted(files):
            if fname.lower().endswith(".ifc"):
                results.append({
                    "path":     os.path.join(root, fname),
                    "filename": fname,
                    "selected": True,
                })
    return results


def get_classification_config() -> list[dict]:
    """Lee ifc_classification_config de PocketBase.
    Retorna lista de records con id, property_name, property_value.
    """
    return pb_client.get_records(
        "ifc_classification_config", params={"perPage": 100}, fallback=[]
    )


def classify_ifc(ifc_path: str, config: dict | None = None) -> tuple[bool, str, int]:
    """Aplica SCFClass_Funciones a todos los IfcProduct con ADIF_00_Codigo_Producto.
    - Si el elemento ya tiene exactamente 1 clasificación correcta → se omite.
    - Si tiene duplicados → se eliminan y se reaplica (corrección).
    - Si no tiene ninguna → se aplica nueva.
    Sobreescribe el archivo original.
    Retorna (ok, mensaje_detallado, elementos_nuevos+corregidos).
    """
    cfg = config or {}
    try:
        model = ifcopenshell.open(ifc_path)

        # En IFC4X3 el atributo URL se llama "Specification"; en IFC4 se llama "Location"
        schema = model.schema.upper()
        url_attr = "Specification" if "IFC4X3" in schema else "Location"

        cls_name = cfg.get("name", "SCFClass_Funciones")

        cls_kwargs: dict = {
            "Name":    cls_name,
            "Source":  cfg.get("source",  "Railway Innovation Hub"),
            "Edition": cfg.get("edition", "V02 2022"),
        }
        edition_date = cfg.get("edition_date", "")
        if edition_date:
            cls_kwargs["EditionDate"] = edition_date
        location = cfg.get("location", "")
        if location:
            cls_kwargs[url_attr] = location

        ref_location = cfg.get("reference_location", "")

        # La entidad IfcClassification se crea solo si hace falta (lazy)
        _classification = None

        def get_classification():
            nonlocal _classification
            if _classification is None:
                _classification = model.create_entity("IfcClassification", **cls_kwargs)
            return _classification

        def get_prop(element, pset_name, prop_name):
            if not hasattr(element, "IsDefinedBy"):
                return None
            for rel in element.IsDefinedBy:
                if not rel.is_a("IfcRelDefinesByProperties"):
                    continue
                pset = rel.RelatingPropertyDefinition
                if not pset.is_a("IfcPropertySet") or pset.Name != pset_name:
                    continue
                for prop in pset.HasProperties:
                    if prop.Name == prop_name:
                        try:
                            return prop.NominalValue.wrappedValue
                        except Exception:
                            return None
            return None

        def get_existing_cls_rels(element):
            """Devuelve lista de IfcRelAssociatesClassification vinculadas a cls_name."""
            rels = []
            if not hasattr(element, "HasAssociations"):
                return rels
            for assoc in element.HasAssociations:
                if not assoc.is_a("IfcRelAssociatesClassification"):
                    continue
                relating = assoc.RelatingClassification
                if relating.is_a("IfcClassification"):
                    if relating.Name == cls_name:
                        rels.append(assoc)
                elif relating.is_a("IfcClassificationReference"):
                    src = relating.ReferencedSource
                    if src and src.is_a("IfcClassification") and src.Name == cls_name:
                        rels.append(assoc)
            return rels

        count_new       = 0
        count_skipped   = 0
        count_corrected = 0

        for element in model.by_type("IfcProduct"):
            codigo = get_prop(element, "ADIF_00_Generales", "ADIF_00_Codigo_Producto")
            if not codigo:
                continue

            existing = get_existing_cls_rels(element)

            if len(existing) == 1:
                # Ya clasificado correctamente → omitir
                count_skipped += 1
                continue

            if len(existing) > 1:
                # Duplicados → eliminar todos y reclasificar
                for rel in existing:
                    ref_entity = rel.RelatingClassification
                    model.remove(rel)
                    try:
                        model.remove(ref_entity)
                    except Exception:
                        pass
                count_corrected += 1
            else:
                count_new += 1

            # Aplicar clasificación (elemento nuevo o corregido)
            desc = get_prop(element, "ADIF_00_Generales", "ADIF_00_Descripcion_Producto")
            try:
                ref_kwargs: dict = {
                    "Identification": str(codigo),
                    "Name": str(desc) if desc else str(codigo),
                    "ReferencedSource": get_classification(),
                }
                # IfcClassificationReference hereda de IfcExternalReference
                # que siempre usa "Location" (en todas las versiones IFC)
                if ref_location:
                    ref_kwargs["Location"] = ref_location
                ref = model.create_entity("IfcClassificationReference", **ref_kwargs)
                model.create_entity(
                    "IfcRelAssociatesClassification",
                    GlobalId=ifcopenshell.guid.new(),
                    RelatedObjects=[element],
                    RelatingClassification=ref,
                )
            except Exception:
                # Si falla la escritura, descontar el contador
                if len(existing) > 1:
                    count_corrected -= 1
                else:
                    count_new -= 1

        atomic_write_model(model, ifc_path)

        parts = []
        if count_new > 0:
            parts.append(f"{count_new} nuevos")
        if count_skipped > 0:
            parts.append(f"{count_skipped} ya clasificados")
        if count_corrected > 0:
            parts.append(f"{count_corrected} duplicados corregidos")
        msg = " | ".join(parts) if parts else "Sin elementos con codigo producto"
        total = count_new + count_corrected
        return True, msg, total
    except Exception as e:
        return False, str(e), 0


def update_element_name(ifc_path: str, step_id: int, new_name: str) -> tuple[bool, str]:
    """Actualiza el atributo Name de un elemento IFC y guarda el archivo."""
    model = _open_cached(ifc_path)
    element = model.by_id(step_id)
    if element is None:
        return False, "Elemento no encontrado"
    try:
        element.Name = new_name
        atomic_write_model(model, ifc_path)
        return True, "OK"
    except Exception as e:
        return False, str(e)


def update_pset_property(ifc_path: str, step_id: int,
                          pset_name: str, prop_name: str,
                          new_value: str) -> tuple[bool, str]:
    """Actualiza el valor de una propiedad de un Pset en el elemento IFC dado.
    Escribe directamente el archivo original. Usa el modelo en caché.
    """
    model = _open_cached(ifc_path)
    element = model.by_id(step_id)
    if element is None:
        return False, "Elemento no encontrado"

    for rel in getattr(element, "IsDefinedBy", []):
        if not rel.is_a("IfcRelDefinesByProperties"):
            continue
        pset = rel.RelatingPropertyDefinition
        if not (pset.is_a("IfcPropertySet") and pset.Name == pset_name):
            continue
        for prop in pset.HasProperties:
            if prop.Name == prop_name:
                try:
                    # Coerción al tipo actual (int/float/bool/str) para no
                    # corromper el tipado del IFC — la UI manda strings.
                    prop.NominalValue.wrappedValue = coerce_property_value(prop, new_value)
                    atomic_write_model(model, ifc_path)
                    return True, "OK"
                except Exception as e:
                    return False, str(e)
        return False, f"Propiedad '{prop_name}' no encontrada en '{pset_name}'"

    return False, f"Pset '{pset_name}' no encontrado en el elemento"


# ── IFC Value Search ─────────────────────────────────────────────────────────

_pset_index_cache: dict[str, dict] = {}


def _build_pset_index(ifc_path: str) -> dict:
    """Construye índice completo: {pset_name: {prop_name: {value_str: [step_ids]}}}.
    Cacheado a nivel de módulo junto con el modelo."""
    if ifc_path in _pset_index_cache:
        return _pset_index_cache[ifc_path]

    model = _open_cached(ifc_path)
    index: dict = {}

    for rel in model.by_type("IfcRelDefinesByProperties"):
        pset = rel.RelatingPropertyDefinition
        if not pset.is_a("IfcPropertySet"):
            continue
        pset_name = pset.Name or "(sin nombre)"
        if pset_name not in index:
            index[pset_name] = {}
        for prop in pset.HasProperties:
            if not prop.is_a("IfcPropertySingleValue"):
                continue
            prop_name = prop.Name or "(sin nombre)"
            nv = prop.NominalValue
            value_str = str(nv.wrappedValue) if nv is not None else ""
            if prop_name not in index[pset_name]:
                index[pset_name][prop_name] = {}
            if value_str not in index[pset_name][prop_name]:
                index[pset_name][prop_name][value_str] = []
            for obj in rel.RelatedObjects:
                index[pset_name][prop_name][value_str].append(obj.id())

    _pset_index_cache[ifc_path] = index
    return index


def _invalidate_pset_index(ifc_path: str) -> None:
    _pset_index_cache.pop(ifc_path, None)


def get_pset_rows(ifc_path: str) -> list[dict]:
    """Devuelve [{name, prop_count}] ordenado alfabéticamente."""
    idx = _build_pset_index(ifc_path)
    return sorted(
        [{"name": pn, "prop_count": len(props)} for pn, props in idx.items()],
        key=lambda x: x["name"],
    )


def get_prop_rows(ifc_path: str, pset_name: str) -> list[dict]:
    """Devuelve [{name, unique_values, element_count}] para el pset dado."""
    idx = _build_pset_index(ifc_path)
    props = idx.get(pset_name, {})
    rows = []
    for prop_name, values in props.items():
        element_count = sum(len(ids) for ids in values.values())
        rows.append({
            "name":          prop_name,
            "unique_values": len(values),
            "element_count": element_count,
        })
    return sorted(rows, key=lambda x: x["name"])


def get_value_rows(ifc_path: str, pset_name: str, prop_name: str) -> list[dict]:
    """Devuelve [{value, count}] ordenado por count desc."""
    idx = _build_pset_index(ifc_path)
    values = idx.get(pset_name, {}).get(prop_name, {})
    rows = [{"value": v, "count": len(ids)} for v, ids in values.items()]
    return sorted(rows, key=lambda x: (-x["count"], x["value"]))


def bulk_update_pset_property(
    ifc_path: str,
    pset_name: str,
    prop_name: str,
    old_value: str,
    new_value: str,
) -> tuple[bool, str, int]:
    """Reemplaza old_value → new_value en TODOS los elementos que tengan ese valor
    en pset_name.prop_name. Actualiza psets únicos una sola vez (por si son compartidos).
    Retorna (ok, mensaje, elementos_actualizados).
    """
    model = _open_cached(ifc_path)
    updated_pset_ids: set[int] = set()
    element_count = 0

    for rel in model.by_type("IfcRelDefinesByProperties"):
        pset = rel.RelatingPropertyDefinition
        if not (pset.is_a("IfcPropertySet") and pset.Name == pset_name):
            continue
        for prop in pset.HasProperties:
            if prop.Name != prop_name:
                continue
            nv = prop.NominalValue
            if nv is None:
                break
            if str(nv.wrappedValue) != old_value:
                break
            # Actualizar la instancia del pset solo una vez
            if pset.id() not in updated_pset_ids:
                try:
                    prop.NominalValue.wrappedValue = coerce_property_value(prop, new_value)
                    updated_pset_ids.add(pset.id())
                except Exception as e:
                    return False, str(e), 0
            # Contar los elementos vinculados a este pset
            element_count += len(list(rel.RelatedObjects))
            break

    if element_count > 0:
        atomic_write_model(model, ifc_path)
        _invalidate_pset_index(ifc_path)
        return True, f"{element_count} elementos actualizados", element_count

    return False, f"No se encontraron elementos con '{old_value}' en {pset_name} → {prop_name}", 0


# ── IFC Material ──────────────────────────────────────────────────────────────

_MAT_COLOR_PALETTE = [
    "#00f0ff", "#ff00ea", "#39ff14", "#ff9900", "#aa88ff",
    "#ff3300", "#ffff00", "#00ccff", "#ff6699", "#66ff99",
    "#ff6600", "#9966ff", "#00ff99", "#ff0066", "#ffcc00",
    "#33ccff", "#ff33cc", "#99ff33", "#ff9933", "#cc99ff",
]


def _mat_color(name: str) -> str:
    """Color neon determinístico a partir del nombre del material."""
    return _MAT_COLOR_PALETTE[abs(hash(name)) % len(_MAT_COLOR_PALETTE)]


def _extract_mat_names(relating_material) -> list[str]:
    """Extrae todos los nombres de material de cualquier tipo de asignación IFC."""
    rm = relating_material
    names: list[str] = []
    try:
        if rm.is_a("IfcMaterial"):
            if rm.Name:
                names.append(rm.Name)
        elif rm.is_a("IfcMaterialList"):
            for m in (rm.Materials or []):
                if m and m.Name:
                    names.append(m.Name)
        elif rm.is_a("IfcMaterialLayerSetUsage"):
            ls = rm.ForLayerSet
            if ls:
                for layer in (ls.MaterialLayers or []):
                    if layer.Material and layer.Material.Name:
                        names.append(layer.Material.Name)
        elif rm.is_a("IfcMaterialLayerSet"):
            for layer in (rm.MaterialLayers or []):
                if layer.Material and layer.Material.Name:
                    names.append(layer.Material.Name)
        elif rm.is_a("IfcMaterialLayer"):
            if rm.Material and rm.Material.Name:
                names.append(rm.Material.Name)
        elif rm.is_a("IfcMaterialConstituentSet"):
            for c in (rm.MaterialConstituents or []):
                if c.Material and c.Material.Name:
                    names.append(c.Material.Name)
        elif rm.is_a("IfcMaterialConstituent"):
            if rm.Material and rm.Material.Name:
                names.append(rm.Material.Name)
        elif rm.is_a("IfcMaterialProfileSetUsage"):
            ps = rm.ForProfileSet
            if ps:
                for prof in (ps.MaterialProfiles or []):
                    if prof.Material and prof.Material.Name:
                        names.append(prof.Material.Name)
        elif rm.is_a("IfcMaterialProfileSet"):
            for prof in (rm.MaterialProfiles or []):
                if prof.Material and prof.Material.Name:
                    names.append(prof.Material.Name)
    except Exception as exc:
        # Degradación: se devuelven los nombres recogidos hasta el fallo
        logger.debug("_extract_mat_names: asignación de material ilegible (%s)", exc)
    return names


def get_material_rows(ifc_path: str) -> list[dict]:
    """Extrae todos los materiales del IFC con estadísticas de uso.
    Retorna [{name, element_count, entity_types, color}] ordenado por element_count desc.
    """
    model = _open_cached(ifc_path)
    # {name: {element_count, entity_types: set}}
    mats: dict[str, dict] = {}

    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            mat_names = _extract_mat_names(rel.RelatingMaterial)
            related = list(rel.RelatedObjects or [])
            for mat_name in mat_names:
                if mat_name not in mats:
                    mats[mat_name] = {"element_count": 0, "entity_types": set()}
                mats[mat_name]["element_count"] += len(related)
                for obj in related:
                    mats[mat_name]["entity_types"].add(obj.is_a())
        except Exception as exc:
            # Una relación corrupta no tumba el listado completo de materiales
            logger.debug("get_material_rows: relación de material omitida (%s)", exc)

    rows = []
    for name, data in sorted(mats.items(), key=lambda x: -x[1]["element_count"]):
        # Truncar entity_types a las 5 más frecuentes no es posible sin conteo individual,
        # así que simplemente tomamos las primeras 5 ordenadas
        et = sorted(data["entity_types"])[:5]
        rows.append({
            "name": name,
            "element_count": data["element_count"],
            "entity_types": et,
            "entity_types_str": " · ".join(et),
            "color": _mat_color(name),
        })
    return rows


def rename_material(
    ifc_path: str,
    old_name: str,
    new_name: str,
) -> tuple[bool, str, int]:
    """Renombra todas las instancias de IfcMaterial con nombre old_name → new_name.
    Escribe el archivo IFC e invalida la caché.
    Retorna (ok, mensaje, instancias_renombradas).
    """
    model = _open_cached(ifc_path)
    count = 0
    for mat in model.by_type("IfcMaterial"):
        if mat.Name == old_name:
            mat.Name = new_name
            count += 1
    if count > 0:
        atomic_write_model(model, ifc_path)
        model_pool.invalidate(ifc_path)
        _pset_index_cache.pop(ifc_path, None)
        return True, f"{count} instancia(s) renombrada(s)", count
    return False, f"Material '{old_name}' no encontrado en {ifc_path.split('/')[-1]}", 0


def get_material_detail_rows(ifc_path: str, mat_name: str) -> list[dict]:
    """Detalle de un material en un archivo: desglose por tipo de entidad con nombres de ejemplo.

    Retorna lista plana para rx.foreach:
      cabecera:  {is_header: True,  label: filename, count: N, entity_type: "", names_preview: ""}
      entidad:   {is_header: False, label: "IfcWall", count: N, entity_type: "IfcWall", names_preview: "A, B +5 más"}
    """
    model = _open_cached(ifc_path)
    filename = ifc_path.split("/")[-1]
    breakdown: dict[str, dict] = {}

    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            if mat_name not in _extract_mat_names(rel.RelatingMaterial):
                continue
            for obj in (rel.RelatedObjects or []):
                et = obj.is_a()
                name = getattr(obj, "Name", None) or ""
                if et not in breakdown:
                    breakdown[et] = {"count": 0, "names": []}
                breakdown[et]["count"] += 1
                if len(breakdown[et]["names"]) < 8 and name:
                    breakdown[et]["names"].append(name)
        except Exception:
            pass

    if not breakdown:
        return []

    total = sum(v["count"] for v in breakdown.values())
    rows: list[dict] = [{
        "is_header": True,
        "label": filename,
        "count": total,
        "entity_type": "",
        "names_preview": "",
    }]
    for et, data in sorted(breakdown.items(), key=lambda x: -x[1]["count"]):
        shown = data["names"][:6]
        extra = data["count"] - len(shown)
        preview = ", ".join(shown)
        if extra > 0:
            preview += f"  +{extra} más"
        rows.append({
            "is_header": False,
            "label": et,
            "count": data["count"],
            "entity_type": et,
            "names_preview": preview,
        })
    return rows


def get_entity_type_rows(ifc_path: str) -> list[dict]:
    """Tipos de entidad con o sin asignación de material, con conteo de elementos."""
    # Caché en disco PRIMERO: si ya se calculó antes para este mismo
    # (ruta, mtime, tamaño), evita abrir/parsear el IFC entero — el costo
    # dominante en archivos de cientos de MB, aunque el resultado sea texto.
    disk = _material_cache_read(ifc_path, "entity_types")
    if disk is not None:
        _entity_type_rows_cache[ifc_path] = disk
        return disk

    # _open_cached: detecta cambio de mtime e invalida caches derivadas
    model = _open_cached(ifc_path)
    if ifc_path in _entity_type_rows_cache:
        return _entity_type_rows_cache[ifc_path]
    et_counts: dict[str, int] = {}
    # Todos los IfcProduct — incluye los que no tienen IfcRelAssociatesMaterial
    for obj in model.by_type("IfcProduct"):
        et = obj.is_a()
        et_counts[et] = et_counts.get(et, 0) + 1
    rows = [
        {"entity_type": et, "element_count": cnt}
        for et, cnt in sorted(et_counts.items(), key=lambda x: -x[1])
    ]
    _entity_type_rows_cache[ifc_path] = rows
    _material_cache_write(ifc_path, "entity_types", rows)
    return rows


def get_materials_for_types(ifc_path: str, entity_types: list[str]) -> list[dict]:
    """Materiales usados por elementos de los tipos indicados."""
    model = _open_cached(ifc_path)
    et_set = set(entity_types)
    mats: dict[str, int] = {}
    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            related = [o for o in (rel.RelatedObjects or []) if o.is_a() in et_set]
            if not related:
                continue
            for mat_name in _extract_mat_names(rel.RelatingMaterial):
                mats[mat_name] = mats.get(mat_name, 0) + len(related)
        except Exception:
            pass
    return sorted(
        [{"name": k, "element_count": v, "color": _mat_color(k)} for k, v in mats.items()],
        key=lambda x: -x["element_count"],
    )


def rename_material_for_types(
    ifc_path: str,
    mat_name: str,
    new_name: str,
    entity_types: list[str],
) -> tuple[bool, str, int]:
    """Renombra/reasigna un material solo para los tipos de entidad indicados.

    Si la relación IfcRelAssociatesMaterial mezcla tipos objetivo y otros:
      - Los elementos objetivo pasan al nuevo material (o se crea uno nuevo).
      - Los demás conservan el material original.
    Retorna (ok, mensaje, elementos_modificados).
    """
    import ifcopenshell.guid as _guid
    model = _open_cached(ifc_path)
    et_set = set(entity_types)
    new_mat = None
    total_changed = 0

    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            rm = rel.RelatingMaterial
            names = _extract_mat_names(rm)
            # "N/A" es el marcador que la UI muestra cuando la relación existe
            # pero no se le pudo extraer nombre (ver el mismo criterio en
            # get_elements_for_types y en _rename_material_mutate).
            if not (mat_name in names or (not names and mat_name == "N/A")):
                continue
            related = list(rel.RelatedObjects or [])
            targets = [o for o in related if o.is_a() in et_set]
            others  = [o for o in related if o.is_a() not in et_set]
            if not targets:
                continue
            # Obtener o crear IfcMaterial de destino
            if new_mat is None:
                existing = [m for m in model.by_type("IfcMaterial") if m.Name == new_name]
                new_mat = existing[0] if existing else model.create_entity("IfcMaterial", Name=new_name)
            if not others and rm.is_a("IfcMaterial"):
                # Relación pura con tipos objetivo → renombrar en sitio
                rm.Name = new_name
            elif not others:
                # Tipo compuesto (LayerSet, etc.) → reasignar relación completa
                rel.RelatingMaterial = new_mat
            else:
                # Relación mixta → dividir: otros conservan, objetivos van al nuevo
                rel.RelatedObjects = others
                model.create_entity(
                    "IfcRelAssociatesMaterial",
                    GlobalId=_guid.new(),
                    OwnerHistory=None,
                    Name=None,
                    Description=None,
                    RelatedObjects=targets,
                    RelatingMaterial=new_mat,
                )
            total_changed += len(targets)
        except Exception as exc:
            # Resultado parcial: la relación problemática se omite y se continúa
            logger.warning("rename_material_for_types: relación omitida (%s)", exc)

    if total_changed > 0:
        atomic_write_model(model, ifc_path)
        model_pool.invalidate(ifc_path)
        _pset_index_cache.pop(ifc_path, None)
        return True, f"{total_changed} elemento(s) reasignados a '{new_name}'", total_changed
    return False, f"Material '{mat_name}' no encontrado para los tipos seleccionados", 0


def _read_pset_material(model, entity, ifc_path: str = "") -> str | None:
    """Lee la propiedad texto 'Material' de cualquier Pset del elemento.

    Esta es la FUENTE PRIORITARIA de material, no un fallback: es el valor
    calculado por BIMRail en Revit (RevitHelpers.GetElementMaterialName,
    escrito en el parámetro 'ifcMaterialAssociation' por XTS_SyncIFC) — la
    MISMA fuente que usa el botón Export de "Objects Materials" para el
    Excel. El IfcMaterial nativo del IFC (IfcRelAssociatesMaterial) puede en
    cambio venir de geometría de librerías externas importadas en Revit
    (p.ej. figuras humanas con IfcMaterialConstituentSet multi-parte:
    "Worker_Skin", "Worker_Jacket"... más un constituyente de relleno como
    "Navisworks_material" heredado del bloque original) que no representa la
    clasificación real del elemento — confirmado inspeccionando
    000000000001-ACO-MOP-ME-DSA-523140-442_MTE_TEMPORAL.ifc: el Pset
    "Materiales y acabados"."Material" decía "MAT_Obrero" (igual que el
    Excel) mientras el IfcMaterial nativo devolvía "Navisworks_material"
    como primer constituyente de una lista de 19.

    Retorna None si el elemento no tiene esta propiedad en ningún Pset
    (típicamente elementos de un origen que nunca pasó por el Sync IFC de
    BIMRail) — en ese caso el llamador debe recurrir al IfcMaterial nativo.
    Retorna su valor si existe, incluido 'N/A' cuando ese fue el valor
    calculado a propósito (p.ej. símbolos de coordinación sin material
    aplicable). Cachea por (ifc_path, step_id) para evitar llamadas
    repetidas a get_psets()."""
    import ifcopenshell.util.element as _eu
    step_id = entity.id() if entity is not None else -1
    if ifc_path and step_id >= 0:
        path_cache = _pset_mat_cache.setdefault(ifc_path, {})
        if step_id in path_cache:
            return path_cache[step_id]
    result = None
    try:
        psets = _eu.get_psets(entity)
        for props in psets.values():
            val = props.get("Material", None)
            if val is not None:
                sval = str(val).strip()
                if sval:
                    result = sval
                    break
    except Exception:
        pass
    if ifc_path and step_id >= 0:
        _pset_mat_cache.setdefault(ifc_path, {})[step_id] = result
    return result


def get_elements_for_types(ifc_path: str, entity_types: list[str]) -> list[dict]:
    """Elementos (con o sin asignación de material) filtrados por tipo de entidad.
    Retorna [{name, step_id, entity_type, material_name, mat_color}].
    Prioriza el Pset 'Material' (valor calculado por BIMRail/Revit — ver
    _read_pset_material) sobre el IfcMaterial nativo del IFC, que para
    geometría importada de librerías externas puede ser ruido de renderizado
    sin relación con la clasificación real (ver docstring de
    _read_pset_material). Solo cuando el elemento no tiene esa propiedad en
    ningún Pset se usa el IfcRelAssociatesMaterial nativo; si tampoco hay eso,
    'N/A'.
    """
    types_key = "|".join(sorted(entity_types))

    # Caché en disco PRIMERO (misma razón que get_entity_type_rows): el costo
    # es el tamaño en bytes del IFC, no la cantidad de elementos consultados.
    disk = _material_cache_read(ifc_path, "elements", extra=types_key)
    if disk is not None:
        path_cache = _elements_cache.setdefault(ifc_path, {})
        path_cache[types_key] = disk
        return disk

    # _open_cached: detecta cambio de mtime e invalida caches derivadas
    model = _open_cached(ifc_path)
    path_cache = _elements_cache.setdefault(ifc_path, {})
    if types_key in path_cache:
        return path_cache[types_key]
    et_set = set(entity_types)
    result: dict[int, dict] = {}

    # Paso 1: registrar todos los elementos de los tipos pedidos (sin material aún)
    for obj in model.by_type("IfcProduct"):
        if obj.is_a() not in et_set:
            continue
        sid = obj.id()
        name = getattr(obj, "Name", None) or ""
        if not name:
            name = f"{obj.is_a()}#{sid}"
        result[sid] = {
            "name": name,
            "step_id": sid,
            "entity_type": obj.is_a(),
            "material_name": None,   # se rellena en paso 2 o 3
            "mat_color": "",
            "_entity": obj,          # temporal para paso 3
        }

    # Paso 2: Pset "Material" (fuente prioritaria — ver _read_pset_material)
    covered: set[int] = set()
    for sid, row in result.items():
        entity = row.get("_entity")
        pset_mat = _read_pset_material(model, entity, ifc_path) if entity is not None else None
        if pset_mat is not None:
            row["material_name"] = pset_mat
            row["mat_color"] = _mat_color(pset_mat)
            row["material_source"] = "pset"
            covered.add(sid)

    # Paso 3: para los no cubiertos (sin esa propiedad en ningún Pset),
    # recurrir al IfcMaterial nativo (IfcRelAssociatesMaterial)
    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            related = [o for o in (rel.RelatedObjects or [])
                       if o.is_a() in et_set and o.id() not in covered]
            if not related:
                continue
            mat_names = _extract_mat_names(rel.RelatingMaterial)
            mat_name = mat_names[0] if mat_names else "N/A"
            for obj in related:
                sid = obj.id()
                if sid in result:
                    result[sid]["material_name"] = mat_name
                    result[sid]["mat_color"] = _mat_color(mat_name)
                    result[sid]["material_source"] = "ifc_rel"
                    covered.add(sid)
        except Exception:
            pass

    # Paso 4: lo que siga sin material tras ambas fuentes → "N/A"
    for sid, row in result.items():
        row.pop("_entity", None)
        if sid not in covered:
            row["material_name"] = "N/A"
            row["mat_color"] = _mat_color("N/A")
            row["material_source"] = "none"

    rows = sorted(result.values(), key=lambda x: (x["entity_type"], x["name"].lower()))
    # Limitar el caché a 8 combinaciones de tipos por archivo para evitar OOM
    if len(path_cache) >= 8:
        path_cache.pop(next(iter(path_cache)))
    path_cache[types_key] = rows
    _material_cache_write(ifc_path, "elements", rows, extra=types_key)
    return rows


def get_materials_for_elements(ifc_path: str, step_ids: list[int]) -> list[dict]:
    """Materiales usados por los elementos indicados (por step_id).
    Prioriza el Pset 'Material' (ver _read_pset_material / get_elements_for_types
    para el porqué) sobre el IfcMaterial nativo; solo recurre a este último
    para los elementos que no tienen esa propiedad en ningún Pset."""
    model = _open_cached(ifc_path)
    id_set = set(step_ids)
    mats: dict[str, int] = {}
    covered: set[int] = set()

    for sid in id_set:
        try:
            entity = model.by_id(sid)
        except Exception:
            continue
        pset_mat = _read_pset_material(model, entity, ifc_path) if entity is not None else None
        if pset_mat is not None:
            mats[pset_mat] = mats.get(pset_mat, 0) + 1
            covered.add(sid)

    uncovered = id_set - covered
    if uncovered:
        for rel in model.by_type("IfcRelAssociatesMaterial"):
            try:
                related = [o for o in (rel.RelatedObjects or []) if o.id() in uncovered]
                if not related:
                    continue
                mat_names = _extract_mat_names(rel.RelatingMaterial)
                mat_name = mat_names[0] if mat_names else "N/A"
                for o in related:
                    mats[mat_name] = mats.get(mat_name, 0) + 1
                    covered.add(o.id())
            except Exception:
                pass

    for sid in id_set - covered:
        mats["N/A"] = mats.get("N/A", 0) + 1

    return sorted(
        [{"name": k, "element_count": v, "color": _mat_color(k)} for k, v in mats.items()],
        key=lambda x: -x["element_count"],
    )


def rename_material_for_elements(
    ifc_path: str,
    mat_name: str,
    new_name: str,
    step_ids: list[int],
) -> tuple[bool, str, int]:
    """Wrapper de compatibilidad para una única operación de rename.

    Ver `rename_material_ops_for_path` para aplicar VARIAS operaciones sobre
    el mismo archivo con una sola apertura + una sola escritura a disco
    (evita serializar el STEP completo una vez por cada material renombrado)."""
    return rename_material_ops_for_path(ifc_path, [(mat_name, new_name, step_ids)])


def rename_material_ops_for_path(
    ifc_path: str,
    ops: list[tuple[str, str, list[int]]],
) -> tuple[bool, str, int]:
    """Aplica N operaciones de rename de material sobre el MISMO archivo con una
    única transacción: abre/muta el modelo en memoria para cada operación y
    escribe a disco UNA sola vez al final (H-perf: escribir un IFC de cientos
    de MB es el costo dominante; hacerlo una vez por lote en vez de una vez
    por material renombrado es la optimización de mayor impacto).

    ops = [(old_material, new_material, step_ids), ...]. Se aplican en el
    orden recibido sobre el mismo modelo en memoria, así que una operación
    puede encadenarse sobre el resultado de la anterior (p.ej. renombrar A→B
    y luego, sobre esos mismos elementos, B→C).
    Retorna (ok, mensaje, total_elementos_modificados)."""
    with model_pool.lock(ifc_path):
        model = _open_cached(ifc_path)
        total_changed = 0
        not_found: list[str] = []
        for mat_name, new_name, step_ids in ops:
            changed = _rename_material_mutate(model, mat_name, new_name, step_ids)
            total_changed += changed
            if changed == 0:
                not_found.append(mat_name)

        if total_changed > 0:
            atomic_write_model(model, ifc_path)
            model_pool.invalidate(ifc_path)
            _pset_index_cache.pop(ifc_path, None)
            return True, f"{total_changed} elemento(s) reasignados", total_changed

        msg = (
            f"Material(es) no encontrado(s) para los elementos seleccionados: {', '.join(not_found)}"
            if not_found else "Sin cambios"
        )
        return False, msg, 0


def _rename_material_mutate(
    model,
    mat_name: str,
    new_name: str,
    step_ids: list[int],
) -> int:
    """Muta el modelo IFC en memoria (sin escribir a disco): reasigna el
    material solo para los elementos indicados (por step_id).

    Misma lógica de split que rename_material_for_types, pero filtra por id
    en lugar de tipo de entidad. Retorna elementos modificados.
    """
    import ifcopenshell.guid as _guid
    id_set = set(step_ids)
    new_mat = None
    total_changed = 0
    handled_ids: set[int] = set()  # ids actually touched below (native rename or new assignment)

    for rel in model.by_type("IfcRelAssociatesMaterial"):
        try:
            rm = rel.RelatingMaterial
            names = _extract_mat_names(rm)
            # "N/A" es el marcador que la UI muestra cuando una relación de
            # material EXISTE pero no se le pudo extraer ningún nombre (p.ej.
            # un IfcMaterial con Name vacío, o un IfcMaterialLayerSet cuyas
            # capas no tienen Material asignado) — get_elements_for_types usa
            # el mismo criterio (`mat_names[0] if mat_names else "N/A"`) para
            # decidir qué mostrar. Antes esta comparación solo miraba nombres
            # extraídos, así que un elemento mostrado como "N/A" nunca podía
            # coincidir con "N/A" aquí; y como además SÍ tenía una relación,
            # el bloque de más abajo (covered_ids) lo daba por cubierto y
            # tampoco lo tocaba: quedaba irrenombrable a cualquier nombre
            # nuevo ("0 elementos, 1 error" en el hub, siempre).
            if not (mat_name in names or (not names and mat_name == "N/A")):
                continue
            related = list(rel.RelatedObjects or [])
            targets = [o for o in related if o.id() in id_set]
            others  = [o for o in related if o.id() not in id_set]
            if not targets:
                continue
            if new_mat is None:
                existing = [m for m in model.by_type("IfcMaterial") if m.Name == new_name]
                new_mat = existing[0] if existing else model.create_entity("IfcMaterial", Name=new_name)
            if not others and rm.is_a("IfcMaterial"):
                rm.Name = new_name
            elif not others:
                rel.RelatingMaterial = new_mat
            else:
                rel.RelatedObjects = others
                model.create_entity(
                    "IfcRelAssociatesMaterial",
                    GlobalId=_guid.new(),
                    OwnerHistory=None,
                    Name=None,
                    Description=None,
                    RelatedObjects=targets,
                    RelatingMaterial=new_mat,
                )
            total_changed += len(targets)
            handled_ids.update(o.id() for o in targets)
        except Exception:
            pass

    # Caso especial: elementos con material solo en Pset (sin IfcRelAssociatesMaterial)
    # Aplica cuando el mat_name pedido no se encontró en ninguna IfcRelAssociatesMaterial
    # del id_set (p.ej. "ByLayer", "N/A") — busca por Pset fallback.
    if True:
        import ifcopenshell.guid as _guid
        covered_ids: set[int] = set()
        for rel in model.by_type("IfcRelAssociatesMaterial"):
            for obj in (rel.RelatedObjects or []):
                if obj.id() in id_set:
                    covered_ids.add(obj.id())
        # Solo elementos sin IfcRelAssociatesMaterial cuyo Pset Material == mat_name
        unassigned = []
        for sid in id_set:
            if sid in covered_ids:
                continue
            try:
                entity = model.by_id(sid)
                pset_mat = _read_pset_material(model, entity) if entity else None
                if pset_mat is not None and pset_mat == mat_name:
                    unassigned.append(sid)
            except Exception:
                pass
        if unassigned:
            if new_mat is None:
                existing = [m for m in model.by_type("IfcMaterial") if m.Name == new_name]
                new_mat = existing[0] if existing else model.create_entity("IfcMaterial", Name=new_name)
            objs = []
            for sid in unassigned:
                try:
                    objs.append(model.by_id(sid))
                except Exception:
                    pass
            if objs:
                model.create_entity(
                    "IfcRelAssociatesMaterial",
                    GlobalId=_guid.new(),
                    OwnerHistory=None,
                    Name=None,
                    Description=None,
                    RelatedObjects=objs,
                    RelatingMaterial=new_mat,
                )
                total_changed += len(objs)
                handled_ids.update(unassigned)

    # Caso especial 2: el elemento YA tiene una IfcRelAssociatesMaterial (a
    # veces con un nombre distinto, p.ej. ya renombrada nativamente en una
    # pasada anterior), pero su Pset "Material" — la fuente que get_elements_
    # for_types/get_materials_for_elements PRIORIZAN para lo que el usuario ve
    # y selecciona como "old_material" — sigue mostrando mat_name (quedó
    # desincronizado). Antes, estos ids quedaban en covered_ids (por tener
    # ALGUNA relación nativa) y por tanto excluidos también del caso especial
    # de arriba: no coincidían por nombre nativo (ya no es mat_name) ni entraban
    # por el fallback de "sin relación nativa" → 0 elementos renombrados con un
    # error confuso de "material no encontrado", aunque el propio Pset seguía
    # diciendo mat_name. Aquí no tocamos la relación nativa (puede ya ser
    # correcta): solo alineamos el texto del Pset con new_name.
    pset_only = []
    for sid in id_set - handled_ids:
        try:
            entity = model.by_id(sid)
            pset_mat = _read_pset_material(model, entity) if entity else None
            if pset_mat is not None and pset_mat == mat_name:
                pset_only.append(sid)
        except Exception:
            pass
    if pset_only:
        _sync_pset_material_property(model, set(pset_only), new_name)
        total_changed += len(pset_only)
        handled_ids.update(pset_only)

    if total_changed > 0:
        # Sincronizar propiedad texto "Material" en todos los Psets de los elementos
        # afectados para que VALUE SEARCH muestre el mismo valor que IFC MATERIAL.
        _sync_pset_material_property(model, id_set, new_name)
    return total_changed


def _sync_pset_material_property(model, step_id_set: set, new_value: str) -> None:
    """Actualiza la propiedad texto 'Material' en todos los Psets de los elementos
    indicados para mantener consistencia con el objeto IfcMaterial asignado."""
    import ifcopenshell.util.element as _eu
    for sid in step_id_set:
        try:
            entity = model.by_id(sid)
        except Exception:
            continue
        if entity is None:
            continue
        # Recorrer todas las relaciones de propiedades del elemento
        for rel in getattr(entity, "IsDefinedBy", []) or []:
            if not rel.is_a("IfcRelDefinesByProperties"):
                continue
            pset = rel.RelatingPropertyDefinition
            if not pset.is_a("IfcPropertySet"):
                continue
            for prop in pset.HasProperties or []:
                if prop.is_a("IfcPropertySingleValue") and prop.Name == "Material":
                    try:
                        prop.NominalValue = model.create_entity(
                            "IfcLabel", wrappedValue=new_value
                        )
                    except Exception:
                        pass


def sync_all_materials_to_psets(ifc_path: str) -> tuple[int, int]:
    """Wrapper transaccional (H1): toma el lock del archivo durante mutate+write."""
    with model_pool.lock(ifc_path):
        return _sync_all_materials_to_psets_impl(ifc_path)


def _sync_all_materials_to_psets_impl(ifc_path: str) -> tuple[int, int]:
    """Sincroniza IfcRelAssociatesMaterial → propiedad 'Material' en Psets,
    SOLO para los elementos que aún no tienen un valor propio ahí.
    Usa el pool LRU en lugar de re-abrir el archivo, y cubre los 9 tipos de relación
    IFC mediante _extract_mat_names (corrección de bug: antes solo cubría 3 tipos).

    Nunca pisa un Pset "Material" que ya tenga un valor: ese valor es la
    clasificación calculada por BIMRail/Revit (ver _read_pset_material), y el
    IfcMaterial nativo puede ser ruido de renderizado sin relación con ella
    (p.ej. constituent-sets multi-parte de geometría importada cuyo primer
    nombre es un constituyente irrelevante tipo "Navisworks_material" —
    confirmado en 442_MTE_TEMPORAL.ifc, donde el Pset real decía
    "MAT_Obrero"). Antes de este fix, este endpoint pisaba ese valor bueno
    con el nombre nativo — este guard es lo que evita que "Sincronizar"
    vuelva a corromper los Psets ya correctos.
    """
    model = _open_cached(ifc_path)

    mat_to_ids: dict[str, set] = {}
    for rel in model.by_type("IfcRelAssociatesMaterial"):
        mat_ref = getattr(rel, "RelatingMaterial", None)
        if mat_ref is None:
            continue
        names = _extract_mat_names(mat_ref)
        if not names:
            continue
        mat_name = names[0]
        related = getattr(rel, "RelatedObjects", None) or []
        for obj in related:
            mat_to_ids.setdefault(mat_name, set()).add(obj.id())

    if not mat_to_ids:
        return 0, 0

    total_elements = sum(len(s) for s in mat_to_ids.values())
    total_psets = 0

    for mat_name, id_set in mat_to_ids.items():
        for sid in id_set:
            try:
                entity = model.by_id(sid)
            except Exception:
                continue
            if entity is None:
                continue
            for rel in getattr(entity, "IsDefinedBy", []) or []:
                if not rel.is_a("IfcRelDefinesByProperties"):
                    continue
                pset = rel.RelatingPropertyDefinition
                if not pset.is_a("IfcPropertySet"):
                    continue
                for prop in pset.HasProperties or []:
                    if prop.is_a("IfcPropertySingleValue") and prop.Name == "Material":
                        try:
                            existing = getattr(prop.NominalValue, "wrappedValue", None) \
                                if prop.NominalValue else None
                            if existing is not None and str(existing).strip():
                                # Ya tiene un valor calculado (autoritativo) — no pisar.
                                continue
                            prop.NominalValue = model.create_entity(
                                "IfcLabel", wrappedValue=mat_name
                            )
                            total_psets += 1
                        except Exception:
                            pass

    atomic_write_model(model, ifc_path)
    model_pool.invalidate(ifc_path)
    _pset_index_cache.pop(ifc_path, None)
    return total_elements, total_psets


# ── Asignaciones persistentes de material ─────────────────────────────────────
_ASSIGNMENTS_FILE = ".ifc_hub_materials.json"


def _assignments_path(folder: str) -> str:
    return os.path.join(folder, _ASSIGNMENTS_FILE)


def save_material_assignments(folder: str, new_entries: list[dict]) -> None:
    """Guarda/actualiza asignaciones esperadas en el JSON del directorio.

    new_entries: [{filename, step_id, element_name, entity_type, material_expected}]
    """
    import json, datetime
    path = _assignments_path(folder)
    existing: dict[str, dict] = {}
    if os.path.exists(path):
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            for a in data.get("assignments", []):
                existing[f"{a['filename']}::{a['step_id']}"] = a
        except Exception:
            pass
    now = datetime.datetime.utcnow().isoformat()
    for entry in new_entries:
        key = f"{entry['filename']}::{entry['step_id']}"
        existing[key] = {**entry, "assigned_at": now}
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(
                {
                    "version": 1,
                    "folder": folder,
                    "saved_at": now,
                    "count": len(existing),
                    "assignments": list(existing.values()),
                },
                f, indent=2, ensure_ascii=False,
            )
    except Exception as exc:
        logger.warning("save_material_assignments: no se pudo escribir %s (%s)", path, exc)


def load_material_assignments(folder: str) -> list[dict]:
    """Carga las asignaciones guardadas. Retorna [] si no existe el archivo."""
    import json
    path = _assignments_path(folder)
    if not os.path.exists(path):
        return []
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f).get("assignments", [])
    except Exception:
        return []


def assignments_count(folder: str) -> int:
    """Número de asignaciones guardadas en el JSON (sin abrir los IFCs)."""
    import json
    path = _assignments_path(folder)
    if not os.path.exists(path):
        return 0
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f).get("count", 0)
    except Exception:
        return 0


def verify_material_assignments(folder: str) -> list[dict]:
    """Compara asignaciones guardadas con el estado actual de los IFCs.

    Retorna discrepancias:
    [{filename, path, step_id, element_name, entity_type,
      material_expected, material_actual, uid}]
    """
    assignments = load_material_assignments(folder)
    if not assignments:
        return []
    discrepancies = []
    for a in assignments:
        ifc_path = os.path.join(folder, a["filename"])
        if not os.path.exists(ifc_path):
            continue
        try:
            model = _open_cached(ifc_path)
            elem  = model.by_id(int(a["step_id"]))
            if elem is None:
                continue
            # Material actual del elemento
            current_mats: list[str] = []
            for rel in model.by_type("IfcRelAssociatesMaterial"):
                if elem in (rel.RelatedObjects or []):
                    current_mats.extend(_extract_mat_names(rel.RelatingMaterial))
            current_mat = current_mats[0] if current_mats else "Sin Material"
            if current_mat != a["material_expected"]:
                discrepancies.append({
                    "filename":          a["filename"],
                    "path":              ifc_path,
                    "step_id":           int(a["step_id"]),
                    "element_name":      a.get("element_name", ""),
                    "entity_type":       a.get("entity_type", ""),
                    "material_expected": a["material_expected"],
                    "material_actual":   current_mat,
                    "uid":               f"{ifc_path}::{a['step_id']}",
                })
        except Exception:
            pass
    return discrepancies


# ── Excel Export / Import ──────────────────────────────────────────────────────

def export_materials_excel(
    element_rows: list[dict],
    filter_material: str | None = None,
) -> bytes:
    """Genera un Excel plano con columnas: ifc_file | entity_type | step_id | element_name | material_name.

    element_rows: lista de dicts tal como los produce get_elements_for_types, enriquecida con
                  'filename' y 'path' por el estado.
    filter_material: si se indica, solo incluye filas de ese material.
    Retorna bytes del fichero .xlsx.
    """
    import io as _io
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
    BLACK = Font(color="000000", name="Calibri", size=11)
    BOLD  = Font(color="000000", name="Calibri", size=11, bold=True)

    rows = element_rows
    if filter_material:
        rows = [r for r in rows if r.get("material_name") == filter_material]

    # Eliminar filas vacías (sin ifc_file ni name)
    rows = [r for r in rows if r.get("filename") or r.get("path") or r.get("name")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Materiales IFC"

    COLS = ["ifc_file", "entity_type", "step_id", "element_name", "material_name"]

    # Cabecera
    for ci, h in enumerate(COLS, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font  = BOLD
        c.fill  = WHITE

    # Datos
    ri = 2
    for r in rows:
        fname = r.get("filename") or os.path.basename(r.get("path", ""))
        vals = [
            fname,
            r.get("entity_type", ""),
            r.get("step_id", 0),
            r.get("name", ""),
            r.get("material_name", ""),
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = BLACK
            c.fill = WHITE
        ri += 1

    # Anchos de columna
    for ci, w in enumerate([50, 26, 10, 50, 42], 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    ws.freeze_panes = "A2"

    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_excel_material_changes(excel_bytes: bytes, folder: str) -> list[dict]:
    """Lee un Excel exportado por export_materials_excel, abre los IFCs y detecta
    qué elementos tienen un material distinto al que indica el Excel.

    Retorna lista de cambios pendientes:
    [{uid, ifc_file, path, step_id, element_name, entity_type, old_material, new_material}]
    """
    import io as _io
    from collections import defaultdict as _dd
    from openpyxl import load_workbook

    buf = _io.BytesIO(excel_bytes)
    wb  = load_workbook(buf, read_only=True, data_only=True)
    ws  = wb.active

    # Parsear cabecera
    headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]

    def _ci(name: str) -> int:
        try:
            return headers.index(name)
        except ValueError:
            raise ValueError(f"Columna requerida '{name}' no encontrada en el Excel")

    ci_file = _ci("ifc_file")
    ci_step = _ci("step_id")
    ci_name = _ci("element_name")
    ci_type = _ci("entity_type")
    ci_mat  = _ci("material_name")

    # Agrupar por fichero IFC
    file_rows: dict[str, list[tuple]] = _dd(list)
    # {ifc_filename: [(step_id, elem_name, entity_type, new_material), ...]}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        ifc_file  = str(row[ci_file]  or "").strip()
        try:
            step_id = int(row[ci_step] or 0)
        except (TypeError, ValueError):
            continue
        elem_name = str(row[ci_name]  or "").strip()
        ent_type  = str(row[ci_type]  or "").strip()
        new_mat   = str(row[ci_mat]   or "").strip()
        if ifc_file and step_id and new_mat:
            file_rows[ifc_file].append((step_id, elem_name, ent_type, new_mat))

    wb.close()

    # Comparar con estado actual de cada IFC
    changes: list[dict] = []

    for ifc_file, file_row_list in file_rows.items():
        ifc_path = os.path.join(folder, ifc_file)
        if not os.path.isfile(ifc_path):
            continue

        id_set = {r[0] for r in file_row_list}
        try:
            model = _open_cached(ifc_path)
        except Exception:
            continue

        # Construir índice: step_id → material_actual
        current_mats: dict[int, str] = {}
        for rel in model.by_type("IfcRelAssociatesMaterial"):
            try:
                mat_names = _extract_mat_names(rel.RelatingMaterial)
                mat_name  = mat_names[0] if mat_names else "Sin Material"
                for obj in (rel.RelatedObjects or []):
                    if obj.id() in id_set:
                        current_mats[obj.id()] = mat_name
            except Exception:
                pass

        for step_id, elem_name, ent_type, new_mat in file_row_list:
            old_mat = current_mats.get(step_id, "Sin Material")
            if old_mat != new_mat:
                changes.append({
                    "uid":          f"{ifc_file}::{step_id}",
                    "ifc_file":     ifc_file,
                    "path":         ifc_path,
                    "step_id":      step_id,
                    "element_name": elem_name,
                    "entity_type":  ent_type,
                    "old_material": old_mat,
                    "new_material": new_mat,
                })

    return changes


def apply_excel_material_changes(changes: list[dict]) -> int:
    """Aplica los cambios de material agrupados por (ifc_path, old_material, new_material),
    con una sola escritura a disco por archivo (todas las operaciones de un mismo
    ifc_path se aplican sobre el mismo modelo en memoria vía rename_material_ops_for_path).
    Retorna total de elementos modificados.
    """
    from collections import defaultdict as _dd

    # Agrupar primero por path, luego por (old_material, new_material) dentro de cada path
    by_path: dict[str, dict[tuple, list[int]]] = _dd(lambda: _dd(list))
    for c in changes:
        by_path[c["path"]][(c["old_material"], c["new_material"])].append(c["step_id"])

    total = 0
    for path, groups in by_path.items():
        ops = [(old_mat, new_mat, step_ids) for (old_mat, new_mat), step_ids in groups.items()]
        try:
            ok, _msg, cnt = rename_material_ops_for_path(path, ops)
            if ok:
                total += cnt
        except Exception:
            pass

    return total


def apply_to_ifc(ifc_path: str, native_name: str, params: list[dict]) -> tuple[bool, str]:
    """Aplica los parámetros ADIF al IfcProject del archivo IFC.
    - Atributos directos: compara y aplica solo si cambiaron.
    - Psets: si no existe → crea (nuevo); si existe 1 con mismo valor → omite;
             si existe 1 con valor distinto → actualiza; si existen 2+ → corrige duplicados.
    Retorna (ok, mensaje_detallado).
    """
    try:
        model = ifcopenshell.open(ifc_path)
        project = model.by_type("IfcProject")[0]
        site = (model.by_type("IfcSite") or [None])[0]

        # Separar parámetros por tipo
        project_attrs: dict[str, str] = {}
        site_attrs:    dict[str, str] = {}
        psets_data:    dict[str, dict[str, str]] = {}

        for row in params:
            p_name = row.get("property_name", "")
            p_val  = str(row.get("property_value", "")).replace("{filename}", native_name)
            group  = row.get("pset_name", "ADIF_Otros")

            if group == "IfcProject (Atributos)":
                project_attrs[p_name] = p_val
            elif group == "IfcSite (Atributos)":
                site_attrs[p_name] = p_val
            else:
                psets_data.setdefault(group, {})[p_name] = p_val

        # ── Helpers ────────────────────────────────────────────────────────
        def find_psets(obj, pset_name: str) -> list:
            """Devuelve lista de (rel, pset) con ese nombre en el objeto."""
            found = []
            if not hasattr(obj, "IsDefinedBy"):
                return found
            for rel in obj.IsDefinedBy:
                if rel.is_a("IfcRelDefinesByProperties"):
                    pset = rel.RelatingPropertyDefinition
                    if pset.is_a("IfcPropertySet") and pset.Name == pset_name:
                        found.append((rel, pset))
            return found

        def pset_values(pset) -> dict[str, str]:
            """Lee los valores actuales de un pset como dict str→str."""
            result = {}
            for prop in pset.HasProperties:
                try:
                    result[prop.Name] = str(prop.NominalValue.wrappedValue)
                except Exception:
                    result[prop.Name] = ""
            return result

        def remove_psets(pset_list: list):
            for rel, pset in pset_list:
                model.remove(rel)
                model.remove(pset)

        def write_pset(obj, g_name: str, props: dict):
            new_pset = ifcopenshell.api.run("pset.add_pset", model, product=obj, name=g_name)
            ifcopenshell.api.run("pset.edit_pset", model, pset=new_pset, properties=props)

        # ── Atributos directos de IfcProject ───────────────────────────────
        attrs_changed = 0
        for p_name, p_val in project_attrs.items():
            current = str(getattr(project, p_name, "") or "")
            if current == p_val:
                continue
            attrs_changed += 1
            if p_name == "Name":          project.Name        = p_val
            elif p_name == "Description": project.Description = p_val
            elif p_name == "ObjectType":  project.ObjectType  = p_val
            elif p_name == "LongName":
                try:    project.LongName = p_val
                except: project.Description = f"{project.Description} | {p_val}"
            elif p_name == "Phase":
                try:    project.Phase = p_val
                except: pass

        # ── Atributos directos de IfcSite ──────────────────────────────────
        if site:
            for p_name, p_val in site_attrs.items():
                current = str(getattr(site, p_name, "") or "")
                if current == p_val:
                    continue
                attrs_changed += 1
                if p_name == "Name":          site.Name        = p_val
                elif p_name == "Description": site.Description = p_val

        # ── Psets ──────────────────────────────────────────────────────────
        count_new        = 0
        count_unchanged  = 0
        count_updated    = 0
        count_corrected  = 0

        for g_name, props in psets_data.items():
            existing = find_psets(project, g_name)

            if len(existing) == 0:
                # No existe → crear
                write_pset(project, g_name, props)
                count_new += 1

            elif len(existing) == 1:
                rel, pset = existing[0]
                current_vals = pset_values(pset)
                if current_vals == {k: str(v) for k, v in props.items()}:
                    # Valores idénticos → omitir
                    count_unchanged += 1
                else:
                    # Valores distintos → actualizar
                    remove_psets(existing)
                    write_pset(project, g_name, props)
                    count_updated += 1

            else:
                # Duplicados → eliminar todos y recrear uno limpio
                remove_psets(existing)
                write_pset(project, g_name, props)
                count_corrected += 1

        needs_write = (attrs_changed + count_new + count_updated + count_corrected) > 0
        if not needs_write and count_unchanged > 0:
            return True, f"SIN CAMBIOS ({count_unchanged} psets ya sincronizados)"

        atomic_write_model(model, ifc_path)

        parts = []
        if attrs_changed > 0:
            parts.append(f"{attrs_changed} atributos actualizados")
        if count_new > 0:
            parts.append(f"{count_new} psets nuevos")
        if count_updated > 0:
            parts.append(f"{count_updated} psets actualizados")
        if count_unchanged > 0:
            parts.append(f"{count_unchanged} psets sin cambios")
        if count_corrected > 0:
            parts.append(f"{count_corrected} duplicados corregidos")
        return True, " | ".join(parts) if parts else "SYNCED"

    except Exception as e:
        return False, str(e)


# ── IFC Nomenclature ──────────────────────────────────────────────────────────

def get_all_entity_type_rows(ifc_path: str) -> list[dict]:
    """Todos los tipos de entidad presentes en el archivo IFC (IfcProduct),
    con conteo de elementos. Usado por IFC Nomenclature (no requiere material)."""
    model = _open_cached(ifc_path)
    et_counts: dict[str, int] = {}
    for elem in model.by_type("IfcProduct"):
        et = elem.is_a()
        et_counts[et] = et_counts.get(et, 0) + 1
    return [
        {"entity_type": et, "element_count": cnt}
        for et, cnt in sorted(et_counts.items(), key=lambda x: -x[1])
    ]


def get_all_elements_for_nomenclature(ifc_path: str) -> list[dict]:
    """Todos los elementos IfcProduct de un archivo IFC.
    Retorna [{step_id, entity_type, name, uid}] sin necesitar filtro de tipo.
    Usado por IFC Nomenclature para validar nomenclatura ADIF sin panel de tipos."""
    model = _open_cached(ifc_path)
    rows = []
    for elem in model.by_type("IfcProduct"):
        name = getattr(elem, "Name", None) or ""
        rows.append({
            "step_id": elem.id(),
            "entity_type": elem.is_a(),
            "name": name,
            "uid": f"{ifc_path}::{elem.id()}",
        })
    rows.sort(key=lambda x: (x["entity_type"], x["name"].lower()))
    return rows


def get_name_rows(element_rows: list[dict]) -> list[dict]:
    """Aggregates unique element names with counts from element_rows.
    Returns [{name, element_count, file_count}] sorted by element_count desc."""
    from collections import defaultdict
    counts = defaultdict(int)
    files  = defaultdict(set)
    for r in element_rows:
        n = r.get("name", "") or ""
        counts[n] += 1
        files[n].add(r.get("filename") or r.get("path", ""))
    result = []
    for name, cnt in sorted(counts.items(), key=lambda x: -x[1]):
        result.append({"name": name, "element_count": cnt, "file_count": len(files[name])})
    return result


def rename_element_name_in_files(
    paths_stepids: dict,   # {ifc_path: [step_id, ...]}
    old_name: str,
    new_name: str,
) -> int:
    """Changes element.Name from old_name to new_name for the given step_ids per file.
    Returns total elements renamed."""
    total = 0
    for path, step_ids in paths_stepids.items():
        total += rename_element_name_single(path, step_ids, old_name, new_name)
    return total


def rename_element_name_single(path: str, step_ids: list, old_name: str, new_name: str) -> int:
    """Renombra elementos en UN archivo IFC. Diseñado para run_in_executor."""
    import ifcopenshell
    try:
        model = ifcopenshell.open(path)
        changed = 0
        for sid in step_ids:
            try:
                elem = model.by_id(int(sid))
                if elem and elem.Name == old_name:
                    elem.Name = new_name
                    changed += 1
            except Exception as exc:
                logger.debug("rename_element_name_single: step_id %s omitido (%s)", sid, exc)
        if changed:
            atomic_write_model(model, path)
            model_pool.invalidate(path)
            _ifc_cache_mtime.pop(path, None)
        return changed
    except Exception as exc:
        logger.warning("rename_element_name_single(%s) falló: %s", path, exc)
        return 0


def export_nomenclature_excel(
    element_rows: list[dict],
    filter_name: str | None = None,
) -> bytes:
    """Exports flat table: ifc_file | entity_type | step_id | element_name.
    filter_name: if given, only rows with that name."""
    import io as _io
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
    BLACK = Font(color="000000", name="Calibri", size=11)
    BOLD  = Font(color="000000", name="Calibri", size=11, bold=True)

    rows = element_rows
    if filter_name is not None:
        rows = [r for r in rows if (r.get("name", "") or "") == filter_name]
    rows = [r for r in rows if r.get("name") or r.get("filename") or r.get("path")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Nomenclatura IFC"

    COLS = ["ifc_file", "entity_type", "step_id", "element_name"]
    for ci, h in enumerate(COLS, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = BOLD
        c.fill = WHITE

    ri = 2
    for r in rows:
        fname = r.get("filename") or os.path.basename(r.get("path", ""))
        vals = [fname, r.get("entity_type", ""), r.get("step_id", 0), r.get("name", "")]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = BLACK
            c.fill = WHITE
        ri += 1

    for ci, w in enumerate([50, 26, 10, 60], 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w
    ws.freeze_panes = "A2"

    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_nomenclature_changes(excel_bytes: bytes, folder: str) -> list[dict]:
    """Reads Excel, opens IFCs, detects elements whose Name differs from Excel.
    Returns [{uid, ifc_file, path, step_id, entity_type, old_name, new_name}]"""
    import io as _io
    from collections import defaultdict as _dd
    from openpyxl import load_workbook

    buf = _io.BytesIO(excel_bytes)
    wb  = load_workbook(buf, read_only=True, data_only=True)
    ws  = wb.active

    headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]

    def _ci(name):
        try:    return headers.index(name)
        except: raise ValueError(f"Columna requerida '{name}' no encontrada")

    ci_file = _ci("ifc_file")
    ci_step = _ci("step_id")
    ci_type = _ci("entity_type")
    ci_name = _ci("element_name")

    file_rows = _dd(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        ifc_file = str(row[ci_file] or "").strip()
        try:
            step_id = int(row[ci_step] or 0)
        except (TypeError, ValueError):
            continue
        ent_type = str(row[ci_type] or "").strip()
        new_name = str(row[ci_name] or "").strip()
        if ifc_file and step_id:
            file_rows[ifc_file].append((step_id, ent_type, new_name))
    wb.close()

    changes = []
    for ifc_file, rows_list in file_rows.items():
        ifc_path = os.path.join(folder, ifc_file)
        if not os.path.isfile(ifc_path):
            continue
        id_set = {r[0] for r in rows_list}
        try:
            model = _open_cached(ifc_path)
        except Exception:
            continue
        # Build index step_id -> current Name
        current_names = {}
        for elem in model.by_type("IfcProduct"):
            if elem.id() in id_set:
                current_names[elem.id()] = elem.Name or ""
        for step_id, ent_type, new_name in rows_list:
            old_name = current_names.get(step_id, "")
            if old_name != new_name:
                changes.append({
                    "uid":       f"{ifc_file}::{step_id}",
                    "ifc_file":  ifc_file,
                    "path":      ifc_path,
                    "step_id":   step_id,
                    "entity_type": ent_type,
                    "old_name":  old_name,
                    "new_name":  new_name,
                })
    return changes


def apply_nomenclature_changes(changes: list[dict]) -> int:
    """Groups changes by path, applies new names per step_id."""
    from collections import defaultdict as _dd
    by_path = _dd(list)
    for c in changes:
        by_path[c["path"]].append(c)
    total = 0
    for path, path_changes in by_path.items():
        sid_to_new = {c["step_id"]: c["new_name"] for c in path_changes}
        total += apply_nomenclature_changes_single(path, sid_to_new)
    return total


def apply_nomenclature_changes_single(path: str, sid_to_new: dict) -> int:
    """Aplica cambios de nomenclatura en UN archivo IFC. Diseñado para run_in_executor."""
    import ifcopenshell
    try:
        model = ifcopenshell.open(path)
        changed = 0
        for sid, new_name in sid_to_new.items():
            try:
                elem = model.by_id(int(sid))
                if elem:
                    elem.Name = new_name
                    changed += 1
            except Exception as exc:
                logger.debug("apply_nomenclature_changes_single: step_id %s omitido (%s)", sid, exc)
        if changed:
            atomic_write_model(model, path)
            model_pool.invalidate(path)
            _ifc_cache_mtime.pop(path, None)
        return changed
    except Exception as exc:
        logger.warning("apply_nomenclature_changes_single(%s) falló: %s", path, exc)
        return 0


# ── IFC Name Type+ID ──────────────────────────────────────────────────────────

def _get_ifc_property(elem, prop_name: str) -> str:
    """Busca un parámetro por nombre en todos los Psets del elemento IFC."""
    try:
        for rel in getattr(elem, "IsDefinedBy", []):
            pset = getattr(rel, "RelatingPropertyDefinition", None)
            if pset is None:
                continue
            for prop in getattr(pset, "HasProperties", []) or []:
                if getattr(prop, "Name", "") == prop_name:
                    nv = getattr(prop, "NominalValue", None)
                    if nv is not None:
                        return str(getattr(nv, "wrappedValue", "")).strip()
    except Exception:
        pass
    return ""


def get_elements_with_adif_id(ifc_path: str) -> list[dict]:
    """Carga todos los IfcProduct con su parámetro ADIF_00_Id_Objeto.
    Retorna [{step_id, entity_type, current_name, id_objeto, new_name,
              has_change, applied}].
    - has_change=True  → tiene ID y el nombre no termina aún en _<ID>.
    - has_change=False → sin ID o ya tiene el ID al final (applied=True).
    """
    model = _open_cached(ifc_path)
    rows = []
    for elem in model.by_type("IfcProduct"):
        current_name = getattr(elem, "Name", None) or ""
        id_obj = _get_ifc_property(elem, "ADIF_00_Id_Objeto")
        if id_obj:
            suffix = f"_{id_obj}"
            if current_name.endswith(suffix):
                new_name   = current_name
                has_change = False
                applied    = True
            else:
                new_name   = f"{current_name}{suffix}"
                has_change = True
                applied    = False
        else:
            new_name   = current_name
            has_change = False
            applied    = False
        rows.append({
            "step_id":      elem.id(),
            "entity_type":  elem.is_a(),
            "current_name": current_name,
            "id_objeto":    id_obj,
            "new_name":     new_name,
            "has_change":   has_change,
            "applied":      applied,
        })
    rows.sort(key=lambda x: (x["current_name"].lower(), x["id_objeto"].lower()))
    return rows


def apply_name_id_renames(changes: list[dict]) -> int:
    """Aplica renombrados name+ID agrupados por archivo IFC.
    changes = [{path, step_id, new_name}]
    Retorna total de elementos renombrados."""
    from collections import defaultdict as _dd
    import ifcopenshell
    by_path: dict = _dd(list)
    for c in changes:
        by_path[c["path"]].append(c)
    total = 0
    for path, items in by_path.items():
        total += apply_name_id_renames_single(path, items)
    return total


def apply_name_id_renames_single(path: str, items: list[dict]) -> int:
    """Aplica renombrados name+ID a UN archivo IFC. Devuelve nº de cambios.
    items = [{step_id, new_name}]
    Diseñado para llamarse desde run_in_executor (no bloquea el event loop)."""
    import ifcopenshell
    try:
        model = ifcopenshell.open(path)
        changed = 0
        for item in items:
            try:
                elem = model.by_id(int(item["step_id"]))
                if elem and hasattr(elem, "Name"):
                    elem.Name = item["new_name"]
                    changed += 1
            except Exception as exc:
                logger.debug("apply_name_id_renames_single: step_id %s omitido (%s)",
                             item.get("step_id"), exc)
        if changed:
            atomic_write_model(model, path)
            model_pool.invalidate(path)
            _ifc_cache_mtime.pop(path, None)
        return changed
    except Exception as exc:
        logger.warning("apply_name_id_renames_single(%s) falló: %s", path, exc)
        return 0
