"""Exportadores Excel (multi-hoja) y PDF (informe ejecutivo) del audit.

Diseñados para no lanzar: ante cualquier error devuelven {"ok": False, "error": ...}.
"""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from .common import (
    STANDARDS, STANDARD_LABELS,
    SEV_INFO, SEV_WARN, SEV_ERROR,
)


# ── Excel ────────────────────────────────────────────────────────────────────

def to_excel(out_path: str, per_file: list[dict], aggregate: dict, folder: str) -> dict:
    """Genera el .xlsx multi-hoja. Requiere openpyxl (ya en requirements)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception as e:
        return {"ok": False, "error": f"openpyxl no disponible: {e}"}

    try:
        wb = Workbook()

        # ── Hoja Resumen ────────────────────────────────────────────────────
        ws = wb.active
        ws.title = "Resumen"
        bold = Font(bold=True)
        hdr_fill = PatternFill("solid", fgColor="1E1E2E")
        hdr_font = Font(bold=True, color="00F0FF")

        ws["A1"] = "IFC Standards Audit"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A3"] = f"Carpeta: {folder}"
        ws["A4"] = f"Archivos auditados: {aggregate.get('files', 0)}"
        ws["A6"] = "Score global"
        ws["B6"] = aggregate.get("score_global", 0)
        ws["A6"].font = bold; ws["B6"].font = bold

        ws["A8"] = "Score por norma"
        ws["A8"].font = hdr_font
        row = 9
        for std in STANDARDS:
            ws.cell(row=row, column=1, value=STANDARD_LABELS[std])
            ws.cell(row=row, column=2, value=aggregate.get("score_by_standard", {}).get(std, 0))
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Issues por severidad").font = hdr_font
        row += 1
        sev = aggregate.get("issues_by_severity", {})
        for k in (SEV_ERROR, SEV_WARN, SEV_INFO):
            ws.cell(row=row, column=1, value=k.upper())
            ws.cell(row=row, column=2, value=sev.get(k, 0))
            row += 1

        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 14

        # ── Hojas por norma ─────────────────────────────────────────────────
        for std in STANDARDS:
            sh = wb.create_sheet(STANDARD_LABELS[std][:30])
            headers = ["archivo", "rule_id", "título", "severidad",
                       "passed", "count", "GlobalIds", "recomendación",
                       "origen", "detalle"]
            for c, h in enumerate(headers, 1):
                cell = sh.cell(row=1, column=c, value=h)
                cell.font = hdr_font
                cell.fill = hdr_fill
            r = 2
            for p in per_file:
                for res in p.get("results", []):
                    if res.get("standard") != std:
                        continue
                    sh.cell(row=r, column=1, value=p.get("file", ""))
                    sh.cell(row=r, column=2, value=res.get("rule_id", ""))
                    sh.cell(row=r, column=3, value=res.get("title", ""))
                    sh.cell(row=r, column=4, value=res.get("severity", ""))
                    sh.cell(row=r, column=5, value="OK" if res.get("passed") else "FAIL")
                    sh.cell(row=r, column=6, value=res.get("issue_count", 0))
                    sh.cell(row=r, column=7, value=", ".join(res.get("sample_ids", [])[:10]))
                    sh.cell(row=r, column=8, value=res.get("recommendation", ""))
                    sh.cell(row=r, column=9, value=res.get("source_hint", ""))
                    sh.cell(row=r, column=10, value=res.get("detail", ""))
                    r += 1
            for col, w in zip("ABCDEFGHIJ", (28, 18, 40, 10, 8, 8, 50, 60, 10, 50)):
                sh.column_dimensions[col].width = w

        # ── Hoja Plan_Mejora ────────────────────────────────────────────────
        plan = wb.create_sheet("Plan_Mejora")
        plan_headers = ["rank", "archivo", "norma", "rule_id", "título",
                        "severidad", "count", "impacto", "origen", "recomendación"]
        for c, h in enumerate(plan_headers, 1):
            cell = plan.cell(row=1, column=c, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
        r = 2
        for item in aggregate.get("improvement_plan", []):
            plan.cell(row=r, column=1, value=item.get("rank", r - 1))
            plan.cell(row=r, column=2, value=item.get("file", ""))
            plan.cell(row=r, column=3, value=STANDARD_LABELS.get(item.get("standard", ""), item.get("standard", "")))
            plan.cell(row=r, column=4, value=item.get("rule_id", ""))
            plan.cell(row=r, column=5, value=item.get("title", ""))
            plan.cell(row=r, column=6, value=item.get("severity", ""))
            plan.cell(row=r, column=7, value=item.get("count", 0))
            plan.cell(row=r, column=8, value=item.get("impact", 0))
            plan.cell(row=r, column=9, value=item.get("source_hint", ""))
            plan.cell(row=r, column=10, value=item.get("recommendation", ""))
            r += 1
        for col, w in zip("ABCDEFGHIJ", (6, 28, 28, 18, 40, 10, 8, 10, 10, 60)):
            plan.column_dimensions[col].width = w

        wb.save(out_path)
        return {"ok": True, "path": out_path, "error": ""}
    except Exception as e:
        return {"ok": False, "error": f"Excel export error: {type(e).__name__}: {str(e)[:200]}"}


# ── PDF ──────────────────────────────────────────────────────────────────────

def to_pdf(out_path: str, per_file: list[dict], aggregate: dict, folder: str) -> dict:
    """Genera el .pdf ejecutivo. Requiere reportlab (ya en requirements)."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
        )
    except Exception as e:
        return {"ok": False, "error": f"reportlab no disponible: {e}"}

    try:
        styles = getSampleStyleSheet()
        h1 = styles["Heading1"]
        h2 = styles["Heading2"]
        h3 = styles["Heading3"]
        body = styles["BodyText"]
        small = ParagraphStyle("small", parent=body, fontSize=8, leading=10)

        doc = SimpleDocTemplate(out_path, pagesize=A4,
                                leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                                topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        story = []

        # ── Portada ─────────────────────────────────────────────────────────
        story.append(Paragraph("IFC Standards Audit — Informe Ejecutivo", h1))
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"<b>Fecha:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", body))
        story.append(Paragraph(f"<b>Carpeta:</b> {folder}", body))
        story.append(Paragraph(f"<b>Archivos auditados:</b> {aggregate.get('files', 0)}", body))
        story.append(Spacer(1, 12))

        # ── KPI resumen ─────────────────────────────────────────────────────
        story.append(Paragraph("Resultado global", h2))
        rows = [["Score global", f"{aggregate.get('score_global', 0)} / 100"]]
        sev = aggregate.get("issues_by_severity", {})
        rows.append(["Issues Error",  sev.get(SEV_ERROR, 0)])
        rows.append(["Issues Warning", sev.get(SEV_WARN, 0)])
        rows.append(["Issues Info",    sev.get(SEV_INFO, 0)])
        t = Table(rows, colWidths=[8 * cm, 5 * cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E1E2E")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.HexColor("#00F0FF")),
            ("FONTNAME",   (0, 0), (-1, -1), "Helvetica"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ]))
        story.append(t)
        story.append(Spacer(1, 12))

        # ── Score por norma ─────────────────────────────────────────────────
        story.append(Paragraph("Score por norma", h2))
        std_rows = [["Norma", "Score"]]
        for s in STANDARDS:
            std_rows.append([STANDARD_LABELS[s],
                             aggregate.get("score_by_standard", {}).get(s, 0)])
        t = Table(std_rows, colWidths=[10 * cm, 3 * cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E1E2E")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ]))
        story.append(t)
        story.append(PageBreak())

        # ── Plan de mejora ──────────────────────────────────────────────────
        story.append(Paragraph("Plan de mejora prioritario (Top 10)", h2))
        plan_rows = [["#", "Archivo", "Norma", "Issue", "Sev.", "Count", "Origen", "Recomendación"]]
        for item in aggregate.get("improvement_plan", [])[:10]:
            fname = os.path.basename(item.get("file", ""))
            if len(fname) > 30:
                fname = fname[:14] + "…" + fname[-14:]
            plan_rows.append([
                Paragraph(str(item.get("rank", "")), small),
                Paragraph(fname, small),
                Paragraph(STANDARD_LABELS.get(item.get("standard", ""), item.get("standard", "")), small),
                Paragraph(item.get("title", "")[:80], small),
                Paragraph(item.get("severity", "").upper(), small),
                Paragraph(str(item.get("count", 0)), small),
                Paragraph(item.get("source_hint", ""), small),
                Paragraph(item.get("recommendation", "")[:240], small),
            ])
        t = Table(plan_rows, colWidths=[0.6 * cm, 3.5 * cm, 2.2 * cm, 3.5 * cm, 1.0 * cm, 1.0 * cm, 1.2 * cm, 4.5 * cm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E1E2E")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.HexColor("#00F0FF")),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ]))
        story.append(t)
        story.append(PageBreak())

        # ── Detalle por archivo ─────────────────────────────────────────────
        for p in per_file:
            story.append(Paragraph(f"Archivo: {p.get('file', '')}", h2))
            story.append(Paragraph(
                f"Autor detectado: <b>{p.get('author', '')}</b> · Schema: {p.get('schema', '')} · "
                f"Score: <b>{p.get('score_global', 0)}/100</b>",
                body))
            story.append(Spacer(1, 6))

            for std in STANDARDS:
                rs = [r for r in p.get("results", []) if r.get("standard") == std]
                if not rs:
                    continue
                story.append(Paragraph(
                    f"{STANDARD_LABELS[std]} · score {p.get('score_by_standard', {}).get(std, 0)}/100",
                    h3))
                rows = [["Rule", "Sev.", "OK", "Count", "Issue", "Recomendación"]]
                for r in rs:
                    rows.append([
                        r.get("rule_id", ""),
                        r.get("severity", ""),
                        "OK" if r.get("passed") else "FAIL",
                        str(r.get("issue_count", 0)),
                        Paragraph(r.get("title", "")[:100], small),
                        Paragraph(r.get("recommendation", "")[:260], small),
                    ])
                t = Table(rows, colWidths=[2.6 * cm, 1.2 * cm, 1.2 * cm, 1.2 * cm, 4.5 * cm, 6.5 * cm], repeatRows=1)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E1E2E")),
                    ("TEXTCOLOR",  (0, 0), (-1, 0), colors.HexColor("#00F0FF")),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))

            story.append(PageBreak())

        doc.build(story)
        return {"ok": True, "path": out_path, "error": ""}
    except Exception as e:
        return {"ok": False, "error": f"PDF export error: {type(e).__name__}: {str(e)[:200]}"}


def default_filename(folder: str, ext: str) -> str:
    """Sugiere nombre tipo `IFCStandardsAudit_<folder>_YYYYmmdd_HHMM.<ext>`."""
    stem = Path(folder).name or "audit"
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    return f"IFCStandardsAudit_{stem}_{ts}.{ext}"
