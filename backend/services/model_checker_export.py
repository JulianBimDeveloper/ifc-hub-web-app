"""model_checker_export.py — Excel & PDF export for IFC Model Checker results.

Receives flat result_rows (list of dicts with level 0=heading, 1=section,
2=check, 3=per-file) plus KPIs and template metadata.
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any


# ─────────────────────────────────────────────────────────────────────────────
# Shared helpers
# ─────────────────────────────────────────────────────────────────────────────

def _fail_elements(row: dict) -> list[dict]:
    """Return fail_elements_json parsed from a check row (level 2)."""
    raw = row.get("fail_elements_json", [])
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except Exception:
            raw = []
    return raw if isinstance(raw, list) else []


# ─────────────────────────────────────────────────────────────────────────────
# Excel export
# ─────────────────────────────────────────────────────────────────────────────

# Dark palette (hex without #)
BG_DARK  = "0D1117"
BG_CARD  = "161B22"
BG_HEAD  = "0A0A1A"
CYAN     = "00E5FF"
GREEN    = "2EA043"
RED_C    = "F85149"
YELLOW   = "D29922"
GRAY     = "8B949E"
WHITE    = "E6EDF3"
BG_PASS  = "0D2818"
BG_FAIL  = "2D1115"
BG_SKIP  = "1C1C2E"


def _xl_fill(hex_color: str):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=hex_color)


def _xl_font(hex_color: str = WHITE, bold: bool = False,
             size: int = 9, name: str = "Calibri"):
    from openpyxl.styles import Font
    return Font(name=name, color=hex_color, bold=bold, size=size)


def _xl_align(horizontal: str = "left", vertical: str = "center",
               wrap: bool = False):
    from openpyxl.styles import Alignment
    return Alignment(horizontal=horizontal, vertical=vertical,
                     wrap_text=wrap)


def _xl_border_thin():
    from openpyxl.styles import Border, Side
    thin = Side(style="thin", color="2A2A3A")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _xl_merge_write(ws, cell_ref: str, end_ref: str, value,
                    fill_hex: str, font_hex: str = WHITE,
                    bold: bool = False, size: int = 9,
                    align: str = "left"):
    ws.merge_cells(f"{cell_ref}:{end_ref}")
    cell = ws[cell_ref]
    cell.value = value
    cell.fill  = _xl_fill(fill_hex)
    cell.font  = _xl_font(font_hex, bold=bold, size=size)
    cell.alignment = _xl_align(align, "center", wrap=True)


def export_excel(
    rows: list[dict],
    kpis: dict,
    template_name: str,
    ifc_files: list[str],
    mode: str = "summary",
) -> bytes:
    """Build a dark-themed Excel workbook and return raw bytes."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for Excel export. "
            "Install it with: pip install openpyxl"
        ) from exc

    wb = Workbook()

    # ── Sheet 1: Resumen ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Resumen"
    ws1.sheet_view.showGridLines = False

    # Set column widths
    col_widths = [6, 18, 18, 32, 14, 6, 10, 10, 10, 10]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # Row 1: title
    _xl_merge_write(ws1, "A1", "J1",
                    "IFC MODEL CHECKER — INFORME DE RESULTADOS",
                    BG_HEAD, WHITE, bold=True, size=14, align="center")
    ws1.row_dimensions[1].height = 28

    # Row 2: date + template
    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws1["A2"].value = "Fecha:"
    ws1["A2"].fill  = _xl_fill(BG_DARK)
    ws1["A2"].font  = _xl_font(GRAY, size=9)
    ws1["B2"].value = date_str
    ws1["B2"].fill  = _xl_fill(BG_DARK)
    ws1["B2"].font  = _xl_font(WHITE, size=9)
    ws1["D2"].value = "Plantilla:"
    ws1["D2"].fill  = _xl_fill(BG_DARK)
    ws1["D2"].font  = _xl_font(GRAY, size=9)
    ws1["E2"].value = template_name
    ws1["E2"].fill  = _xl_fill(BG_DARK)
    ws1["E2"].font  = _xl_font(WHITE, size=9)
    for col in ["C", "F", "G", "H", "I", "J"]:
        ws1[f"{col}2"].fill = _xl_fill(BG_DARK)
    ws1.row_dimensions[2].height = 16

    # Row 3: IFC files
    files_str = ", ".join(ifc_files[:4])
    if len(ifc_files) > 4:
        files_str += f" (+{len(ifc_files) - 4} más)"
    ws1["A3"].value = "Modelos IFC:"
    ws1["A3"].fill  = _xl_fill(BG_DARK)
    ws1["A3"].font  = _xl_font(GRAY, size=9)
    ws1.merge_cells("B3:J3")
    ws1["B3"].value = files_str
    ws1["B3"].fill  = _xl_fill(BG_DARK)
    ws1["B3"].font  = _xl_font(WHITE, size=9)
    ws1["B3"].alignment = _xl_align("left", "center")
    ws1.row_dimensions[3].height = 16

    # Row 4: spacer
    for col in range(1, 11):
        ws1.cell(row=4, column=col).fill = _xl_fill(BG_DARK)
    ws1.row_dimensions[4].height = 6

    # Row 5: KPI block (5 groups of 2 merged columns)
    total     = kpis.get("total", 0)
    pass_c    = kpis.get("pass_c", 0)
    fail_c    = kpis.get("fail_c", 0)
    skip_c    = kpis.get("skip_c", 0)
    pct       = kpis.get("pct", 0)

    kpi_defs = [
        ("A5", "B5", f"TOTAL\n{total}",        BG_CARD, CYAN),
        ("C5", "D5", f"PASADOS\n{pass_c}",     BG_PASS, GREEN),
        ("E5", "F5", f"FALLADOS\n{fail_c}",    BG_FAIL, RED_C),
        ("G5", "H5", f"N/A\n{skip_c}",         BG_SKIP, GRAY),
        ("I5", "J5", f"CUMPLIMIENTO\n{pct}%",  BG_CARD, CYAN),
    ]
    for start, end, val, bg, fg in kpi_defs:
        _xl_merge_write(ws1, start, end, val, bg, fg,
                        bold=True, size=11, align="center")
    ws1.row_dimensions[5].height = 36

    # Row 6: spacer
    for col in range(1, 11):
        ws1.cell(row=6, column=col).fill = _xl_fill(BG_DARK)
    ws1.row_dimensions[6].height = 6

    # Row 7: table headers
    headers = ["#", "HEADING", "SECCIÓN", "CHECK", "TIPO",
               "REQ", "ESTADO", "PASADOS", "FALLADOS", "TOTAL"]
    from openpyxl.styles import Border, Side
    bottom_thin = Border(bottom=Side(style="thin", color=CYAN))
    for col_idx, hdr in enumerate(headers, 1):
        cell = ws1.cell(row=7, column=col_idx, value=hdr)
        cell.fill      = _xl_fill(BG_HEAD)
        cell.font      = _xl_font(CYAN, bold=True, size=9)
        cell.alignment = _xl_align("center", "center")
        cell.border    = bottom_thin
    ws1.row_dimensions[7].height = 18

    # Row 8+: check rows (level == 2)
    current_heading  = ""
    current_section  = ""
    data_row = 8

    for r in rows:
        lv = r.get("level", -1)
        if lv == 0:
            current_heading = r.get("label", "")
        elif lv == 1:
            current_section = r.get("label", "")
        elif lv == 2:
            status      = r.get("status", "skip")
            is_required = r.get("is_required", False)
            passed      = r.get("passed", 0)
            failed      = r.get("failed", 0)
            total_c     = r.get("total", 0)

            # Row background
            if status == "pass":
                row_bg, text_fg = BG_PASS, WHITE
            elif status == "fail":
                row_bg, text_fg = BG_FAIL, WHITE
            else:
                row_bg, text_fg = BG_SKIP, GRAY

            status_sym = "✓" if status == "pass" else ("✗" if status == "fail" else "—")
            status_fg  = GREEN if status == "pass" else (RED_C if status == "fail" else GRAY)
            bold_row   = is_required and status == "fail"

            values = [
                r.get("number", ""),
                current_heading,
                current_section,
                r.get("label", ""),
                r.get("check_type", ""),
                "SI" if is_required else "",
                status_sym,
                passed,
                failed,
                total_c,
            ]
            border = _xl_border_thin()
            for col_idx, val in enumerate(values, 1):
                cell = ws1.cell(row=data_row, column=col_idx, value=val)
                cell.fill   = _xl_fill(row_bg)
                cell.border = border
                cell.alignment = _xl_align(
                    "center" if col_idx in (1, 6, 7, 8, 9, 10) else "left",
                    "center",
                )
                # Status column special color
                if col_idx == 7:
                    cell.font = _xl_font(status_fg, bold=True, size=10)
                else:
                    cell.font = _xl_font(
                        RED_C if (col_idx == 8 and failed > 0) else text_fg,
                        bold=bold_row,
                        size=9,
                    )
            ws1.row_dimensions[data_row].height = 16
            data_row += 1

    # Footer summary row
    ws1.merge_cells(f"A{data_row}:F{data_row}")
    ws1[f"A{data_row}"].value = "RESUMEN FINAL"
    ws1[f"A{data_row}"].fill  = _xl_fill(BG_HEAD)
    ws1[f"A{data_row}"].font  = _xl_font(CYAN, bold=True, size=9)
    ws1[f"A{data_row}"].alignment = _xl_align("right", "center")
    ws1[f"G{data_row}"].value = f"{pct}%"
    ws1[f"G{data_row}"].fill  = _xl_fill(BG_HEAD)
    ws1[f"G{data_row}"].font  = _xl_font(CYAN, bold=True, size=10)
    ws1[f"G{data_row}"].alignment = _xl_align("center", "center")
    ws1[f"H{data_row}"].value = pass_c
    ws1[f"H{data_row}"].fill  = _xl_fill(BG_HEAD)
    ws1[f"H{data_row}"].font  = _xl_font(GREEN, bold=True, size=9)
    ws1[f"H{data_row}"].alignment = _xl_align("center", "center")
    ws1[f"I{data_row}"].value = fail_c
    ws1[f"I{data_row}"].fill  = _xl_fill(BG_HEAD)
    ws1[f"I{data_row}"].font  = _xl_font(RED_C, bold=True, size=9)
    ws1[f"I{data_row}"].alignment = _xl_align("center", "center")
    ws1[f"J{data_row}"].value = total
    ws1[f"J{data_row}"].fill  = _xl_fill(BG_HEAD)
    ws1[f"J{data_row}"].font  = _xl_font(WHITE, bold=True, size=9)
    ws1[f"J{data_row}"].alignment = _xl_align("center", "center")

    # Freeze panes at A8
    ws1.freeze_panes = "A8"

    # ── Sheet 2: Fallos ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Fallos")
    ws2.sheet_view.showGridLines = False

    fail_col_widths = [28, 24, 12, 28, 20, 24, 10]
    fail_col_names  = ["A", "B", "C", "D", "E", "F", "G"]
    for col_letter, w in zip(fail_col_names, fail_col_widths):
        ws2.column_dimensions[col_letter].width = w

    # Header
    ws2.merge_cells("A1:G1")
    ws2["A1"].value = "ELEMENTOS CON FALLOS"
    ws2["A1"].fill  = _xl_fill(BG_FAIL)
    ws2["A1"].font  = _xl_font(RED_C, bold=True, size=12)
    ws2["A1"].alignment = _xl_align("center", "center")
    ws2.row_dimensions[1].height = 24

    # Column headers
    fail_headers = ["CHECK", "ARCHIVO IFC", "STEP ID",
                    "ELEMENTO", "TIPO IFC", "VALOR ACTUAL", "REQUERIDO"]
    for col_idx, hdr in enumerate(fail_headers, 1):
        cell = ws2.cell(row=2, column=col_idx, value=hdr)
        cell.fill      = _xl_fill(BG_HEAD)
        cell.font      = _xl_font(CYAN, bold=True, size=9)
        cell.alignment = _xl_align("center", "center")
        cell.border    = _xl_border_thin()
    ws2.row_dimensions[2].height = 16

    fail_data_row = 3
    alt = False
    for r in rows:
        if r.get("level") != 2 or r.get("status") != "fail":
            continue
        check_label = f"{r.get('number', '')} {r.get('label', '')}"
        is_required = r.get("is_required", False)
        elements    = _fail_elements(r)
        for elem in elements:
            row_bg = BG_DARK if alt else BG_CARD
            alt    = not alt
            vals   = [
                check_label,
                elem.get("ifc_file", ""),
                elem.get("step_id", ""),
                elem.get("name", ""),
                elem.get("type", ""),
                elem.get("current_value", ""),
                "SI" if is_required else "",
            ]
            border = _xl_border_thin()
            for col_idx, val in enumerate(vals, 1):
                cell = ws2.cell(row=fail_data_row, column=col_idx, value=val)
                cell.fill      = _xl_fill(row_bg)
                cell.font      = _xl_font(WHITE, size=9)
                cell.alignment = _xl_align("left", "center")
                cell.border    = border
            ws2.row_dimensions[fail_data_row].height = 15
            fail_data_row += 1

    if fail_data_row == 3:
        ws2.merge_cells("A3:G3")
        ws2["A3"].value = "Sin elementos con fallos"
        ws2["A3"].fill  = _xl_fill(BG_DARK)
        ws2["A3"].font  = _xl_font(GRAY, size=9)
        ws2["A3"].alignment = _xl_align("center", "center")

    # ── Sheet 3: Por Archivo (detail mode only) ───────────────────────────────
    if mode == "detail":
        ws3 = wb.create_sheet("Por Archivo")
        ws3.sheet_view.showGridLines = False

        pf_col_widths = [12, 24, 28, 10, 10, 10, 10, 28]
        pf_col_letters = [get_column_letter(i) for i in range(1, 9)]
        for col_letter, w in zip(pf_col_letters, pf_col_widths):
            ws3.column_dimensions[col_letter].width = w

        # Header
        ws3.merge_cells("A1:H1")
        ws3["A1"].value = "RESULTADOS POR ARCHIVO"
        ws3["A1"].fill  = _xl_fill(BG_HEAD)
        ws3["A1"].font  = _xl_font(CYAN, bold=True, size=12)
        ws3["A1"].alignment = _xl_align("center", "center")
        ws3.row_dimensions[1].height = 24

        pf_headers = ["#", "ARCHIVO IFC", "CHECK PADRE",
                      "ESTADO", "PASADOS", "FALLADOS", "TOTAL", "INFO"]
        for col_idx, hdr in enumerate(pf_headers, 1):
            cell = ws3.cell(row=2, column=col_idx, value=hdr)
            cell.fill      = _xl_fill(BG_HEAD)
            cell.font      = _xl_font(CYAN, bold=True, size=9)
            cell.alignment = _xl_align("center", "center")
            cell.border    = _xl_border_thin()
        ws3.row_dimensions[2].height = 16

        pf_data_row = 3
        _check_labels: dict[str, str] = {}
        for r in rows:
            if r.get("level") == 2:
                _check_labels[r["row_id"]] = r.get("label", "")

        for r in rows:
            if r.get("level") != 3:
                continue
            status = r.get("status", "skip")
            if status == "pass":
                row_bg, text_fg = BG_PASS, WHITE
            elif status == "fail":
                row_bg, text_fg = BG_FAIL, WHITE
            else:
                row_bg, text_fg = BG_SKIP, GRAY

            status_sym = "✓" if status == "pass" else ("✗" if status == "fail" else "—")
            parent_label = _check_labels.get(r.get("parent_id", ""), "")
            vals = [
                r.get("number", ""),
                r.get("label", ""),
                parent_label,
                status_sym,
                r.get("passed", 0),
                r.get("failed", 0),
                r.get("total", 0),
                r.get("info", ""),
            ]
            border = _xl_border_thin()
            for col_idx, val in enumerate(vals, 1):
                cell = ws3.cell(row=pf_data_row, column=col_idx, value=val)
                cell.fill      = _xl_fill(row_bg)
                cell.font      = _xl_font(text_fg, size=9)
                cell.alignment = _xl_align("center" if col_idx in (1, 4, 5, 6, 7) else "left",
                                           "center")
                cell.border    = border
            ws3.row_dimensions[pf_data_row].height = 15
            pf_data_row += 1

    # Serialize to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ─────────────────────────────────────────────────────────────────────────────
# PDF export
# ─────────────────────────────────────────────────────────────────────────────

def export_pdf(
    rows: list[dict],
    kpis: dict,
    template_name: str,
    ifc_files: list[str],
    mode: str = "summary",
) -> bytes:
    """Build a professional light-themed PDF and return raw bytes."""
    try:
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, PageBreak, HRFlowable,
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.colors import HexColor, white, black, Color
        from reportlab.lib.units import cm
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    except ImportError as exc:
        raise ImportError(
            "reportlab is required for PDF export. "
            "Install it with: pip install reportlab"
        ) from exc

    # Palette
    c_navy  = HexColor("#1a237e")
    c_blue  = HexColor("#1565c0")
    c_green = HexColor("#2e7d32")
    c_red   = HexColor("#c62828")
    c_amber = HexColor("#e65100")
    c_gray  = HexColor("#546e7a")
    c_light = HexColor("#eceff1")
    c_dark  = HexColor("#263238")
    c_head_bg = HexColor("#1565c0")
    c_row_alt = HexColor("#f5f7fa")
    c_pass_bg = HexColor("#e8f5e9")
    c_fail_bg = HexColor("#ffebee")
    c_skip_bg = HexColor("#f5f5f5")
    c_group_bg = HexColor("#e3f2fd")

    buf = io.BytesIO()
    page_w, page_h = A4
    margin = 1.5 * cm

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin + 1.0 * cm,
        bottomMargin=margin + 0.8 * cm,
        title=f"IFC Model Checker — {template_name}",
        author="IFC HUB",
    )

    styles = getSampleStyleSheet()
    body_w = page_w - 2 * margin

    def _style(name: str, **kwargs) -> ParagraphStyle:
        base = styles["Normal"]
        return ParagraphStyle(name, parent=base, **kwargs)

    st_title    = _style("title",    fontSize=18, textColor=white,
                          fontName="Helvetica-Bold", alignment=TA_LEFT)
    st_subtitle = _style("subtitle", fontSize=10, textColor=white,
                          fontName="Helvetica",     alignment=TA_RIGHT)
    st_meta     = _style("meta",     fontSize=9,  textColor=c_gray,
                          fontName="Helvetica")
    st_body     = _style("body",     fontSize=9,  textColor=c_dark,
                          fontName="Helvetica")
    st_bold     = _style("bold",     fontSize=9,  textColor=c_dark,
                          fontName="Helvetica-Bold")
    st_small    = _style("small",    fontSize=8,  textColor=c_gray,
                          fontName="Helvetica")
    st_heading  = _style("hdr",      fontSize=10, textColor=c_navy,
                          fontName="Helvetica-Bold")
    st_section  = _style("sect",     fontSize=9,  textColor=c_blue,
                          fontName="Helvetica-Bold")

    story: list[Any] = []
    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    # ── 1. Header block ───────────────────────────────────────────────────────
    header_table_data = [[
        Paragraph("IFC MODEL CHECKER", st_title),
        Paragraph("INFORME DE RESULTADOS", st_subtitle),
    ]]
    header_table = Table(header_table_data, colWidths=[body_w * 0.6, body_w * 0.4])
    header_table.setStyle(TableStyle([
        ("BACKGROUND",  (0, 0), (-1, -1), c_head_bg),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING",    (0, 0), (-1, -1), 12),
        ("LEFTPADDING",   (0, 0), (0, -1), 14),
        ("RIGHTPADDING",  (-1, 0), (-1, -1), 14),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 0.3 * cm))

    # Meta info
    files_str = ", ".join(ifc_files[:4])
    if len(ifc_files) > 4:
        files_str += f" (+{len(ifc_files) - 4} más)"
    meta_data = [
        [Paragraph(f"<b>Plantilla:</b> {template_name}", st_body),
         Paragraph(f"<b>Fecha:</b> {date_str}", st_body)],
        [Paragraph(f"<b>Modelos IFC:</b> {files_str}", st_body),
         Paragraph(f"<b>Archivos:</b> {len(ifc_files)}", st_body)],
    ]
    meta_table = Table(meta_data, colWidths=[body_w * 0.7, body_w * 0.3])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), c_light),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("LEFTPADDING",   (0, 0), (-1, -1), 8),
        ("GRID",          (0, 0), (-1, -1), 0.5, HexColor("#cfd8dc")),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 0.4 * cm))

    # ── 2. KPI boxes ─────────────────────────────────────────────────────────
    total  = kpis.get("total", 0)
    pass_c = kpis.get("pass_c", 0)
    fail_c = kpis.get("fail_c", 0)
    skip_c = kpis.get("skip_c", 0)
    pct    = kpis.get("pct", 0)

    def _kpi_cell(label: str, val, color) -> list:
        return [
            Paragraph(f'<font size="16"><b>{val}</b></font>', ParagraphStyle(
                "kv", parent=styles["Normal"], alignment=TA_CENTER,
                textColor=white, fontName="Helvetica-Bold")),
            Paragraph(f'<font size="8">{label}</font>', ParagraphStyle(
                "kl", parent=styles["Normal"], alignment=TA_CENTER,
                textColor=white, fontName="Helvetica")),
        ]

    kpi_data = [[
        Paragraph(f'<font size="16"><b>{total}</b></font>',
                  ParagraphStyle("kv0", parent=styles["Normal"],
                                 alignment=TA_CENTER, textColor=white,
                                 fontName="Helvetica-Bold")),
        Paragraph(f'<font size="16"><b>{pass_c}</b></font>',
                  ParagraphStyle("kv1", parent=styles["Normal"],
                                 alignment=TA_CENTER, textColor=white,
                                 fontName="Helvetica-Bold")),
        Paragraph(f'<font size="16"><b>{fail_c}</b></font>',
                  ParagraphStyle("kv2", parent=styles["Normal"],
                                 alignment=TA_CENTER, textColor=white,
                                 fontName="Helvetica-Bold")),
        Paragraph(f'<font size="16"><b>{skip_c}</b></font>',
                  ParagraphStyle("kv3", parent=styles["Normal"],
                                 alignment=TA_CENTER, textColor=white,
                                 fontName="Helvetica-Bold")),
        Paragraph(f'<font size="16"><b>{pct}%</b></font>',
                  ParagraphStyle("kv4", parent=styles["Normal"],
                                 alignment=TA_CENTER, textColor=white,
                                 fontName="Helvetica-Bold")),
    ], [
        Paragraph("TOTAL",        ParagraphStyle("kl0", parent=styles["Normal"],
                                                  alignment=TA_CENTER, textColor=white,
                                                  fontName="Helvetica", fontSize=8)),
        Paragraph("PASADOS",      ParagraphStyle("kl1", parent=styles["Normal"],
                                                  alignment=TA_CENTER, textColor=white,
                                                  fontName="Helvetica", fontSize=8)),
        Paragraph("FALLADOS",     ParagraphStyle("kl2", parent=styles["Normal"],
                                                  alignment=TA_CENTER, textColor=white,
                                                  fontName="Helvetica", fontSize=8)),
        Paragraph("N/A",          ParagraphStyle("kl3", parent=styles["Normal"],
                                                  alignment=TA_CENTER, textColor=white,
                                                  fontName="Helvetica", fontSize=8)),
        Paragraph("CUMPLIMIENTO", ParagraphStyle("kl4", parent=styles["Normal"],
                                                  alignment=TA_CENTER, textColor=white,
                                                  fontName="Helvetica", fontSize=8)),
    ]]

    kpi_w = body_w / 5
    kpi_table = Table(kpi_data, colWidths=[kpi_w] * 5)
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (0, -1), c_blue),
        ("BACKGROUND",    (1, 0), (1, -1), c_green),
        ("BACKGROUND",    (2, 0), (2, -1), c_red),
        ("BACKGROUND",    (3, 0), (3, -1), c_gray),
        ("BACKGROUND",    (4, 0), (4, -1), c_navy),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("LINEAFTER",     (0, 0), (-2, -1), 1, white),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── 3. Results table ─────────────────────────────────────────────────────
    story.append(Paragraph("RESULTADOS DE LA VALIDACIÓN", st_heading))
    story.append(Spacer(1, 0.2 * cm))

    col_w = [1.2 * cm, body_w * 0.33, body_w * 0.15, 1.2 * cm,
             1.6 * cm, 1.4 * cm, 1.4 * cm, 1.4 * cm]
    # Adjust last column to fill remaining space
    remaining = body_w - sum(col_w[:-1])
    col_w[-1] = max(remaining, 1.2 * cm)

    tbl_headers = [
        Paragraph("<b>#</b>",      st_small),
        Paragraph("<b>Check</b>",  st_small),
        Paragraph("<b>Tipo</b>",   st_small),
        Paragraph("<b>Req</b>",    st_small),
        Paragraph("<b>Estado</b>", st_small),
        Paragraph("<b>✓</b>",      st_small),
        Paragraph("<b>✗</b>",      st_small),
        Paragraph("<b>Total</b>",  st_small),
    ]
    tbl_data  = [tbl_headers]
    tbl_style = [
        # Header
        ("BACKGROUND",    (0, 0), (-1, 0), c_navy),
        ("TEXTCOLOR",     (0, 0), (-1, 0), white),
        ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING",    (0, 0), (-1, 0), 6),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("GRID",          (0, 0), (-1, -1), 0.3, HexColor("#b0bec5")),
    ]

    row_idx = 1
    current_heading = ""
    _alt = False
    for r in rows:
        lv = r.get("level", -1)
        if lv == 0:
            heading_label = r.get("label", "")
            if heading_label != current_heading:
                current_heading = heading_label
                tbl_data.append([
                    Paragraph(
                        f'<b>{r.get("number", "")}  {heading_label}</b>',
                        ParagraphStyle("hdr_row", parent=styles["Normal"],
                                       textColor=c_navy, fontName="Helvetica-Bold",
                                       fontSize=9)
                    ),
                    "", "", "", "", "", "", "",
                ])
                tbl_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), c_group_bg))
                tbl_style.append(("SPAN",        (0, row_idx), (-1, row_idx)))
                tbl_style.append(("FONTNAME",    (0, row_idx), (-1, row_idx),
                                   "Helvetica-Bold"))
                tbl_style.append(("LEFTPADDING", (0, row_idx), (-1, row_idx), 6))
                row_idx += 1
        elif lv == 2:
            status      = r.get("status", "skip")
            is_required = r.get("is_required", False)
            passed      = r.get("passed", 0)
            failed      = r.get("failed", 0)
            total_c     = r.get("total", 0)

            if status == "pass":
                row_bg = c_pass_bg
                st_lbl = _style(f"rp{row_idx}", fontSize=8, textColor=c_dark,
                                 fontName="Helvetica")
                status_text = Paragraph('<font color="#2e7d32"><b>PASS</b></font>',
                                         styles["Normal"])
            elif status == "fail":
                row_bg = c_fail_bg
                fn = "Helvetica-Bold" if is_required else "Helvetica"
                st_lbl = _style(f"rp{row_idx}", fontSize=8, textColor=c_dark,
                                 fontName=fn)
                status_text = Paragraph('<font color="#c62828"><b>FALLO</b></font>',
                                         styles["Normal"])
            else:
                row_bg = c_row_alt if _alt else white
                st_lbl = _style(f"rp{row_idx}", fontSize=8, textColor=c_gray,
                                 fontName="Helvetica")
                status_text = Paragraph('<font color="#546e7a">N/A</font>',
                                         styles["Normal"])
            _alt = not _alt

            req_text = Paragraph('<b>SI</b>' if is_required else "",
                                  st_small)
            tbl_data.append([
                Paragraph(r.get("number", ""), st_small),
                Paragraph(r.get("label", ""), st_lbl),
                Paragraph(r.get("check_type", ""), st_small),
                req_text,
                status_text,
                Paragraph(str(passed), st_small),
                Paragraph(str(failed), st_small),
                Paragraph(str(total_c), st_small),
            ])
            tbl_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), row_bg))
            tbl_style.append(("BOTTOMPADDING", (0, row_idx), (-1, row_idx), 3))
            tbl_style.append(("TOPPADDING",    (0, row_idx), (-1, row_idx), 3))
            row_idx += 1

    results_table = Table(tbl_data, colWidths=col_w, repeatRows=1)
    results_table.setStyle(TableStyle(tbl_style))
    story.append(results_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── 4. Failures section ───────────────────────────────────────────────────
    story.append(HRFlowable(width="100%", thickness=1, color=c_red))
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph("ELEMENTOS CON FALLOS", _style(
        "fail_hdr", fontSize=12, textColor=c_red,
        fontName="Helvetica-Bold")))
    story.append(Spacer(1, 0.3 * cm))

    current_group = ""
    any_fail = False

    for r in rows:
        if r.get("level") != 2 or r.get("status") != "fail":
            continue
        any_fail = True
        elements = _fail_elements(r)
        if not elements:
            continue

        # Group header
        group_key = f"{r.get('h_id', '')}_{r.get('s_id', '')}"
        check_label = f"{r.get('number', '')} — {r.get('label', '')}"
        story.append(Paragraph(check_label, st_section))
        story.append(Spacer(1, 0.1 * cm))

        if mode == "detail":
            # Full element table
            fail_headers_pdf = [
                Paragraph("<b>Archivo IFC</b>",  st_small),
                Paragraph("<b>STEP ID</b>",       st_small),
                Paragraph("<b>Elemento</b>",       st_small),
                Paragraph("<b>Tipo IFC</b>",       st_small),
                Paragraph("<b>Valor Actual</b>",   st_small),
            ]
            fail_tbl_data = [fail_headers_pdf]
            fail_tbl_style = [
                ("BACKGROUND",    (0, 0), (-1, 0), c_red),
                ("TEXTCOLOR",     (0, 0), (-1, 0), white),
                ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE",      (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING",    (0, 0), (-1, -1), 3),
                ("GRID",          (0, 0), (-1, -1), 0.3, HexColor("#ffcdd2")),
                ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
            ]
            elem_alt = False
            for e_idx, elem in enumerate(elements, 1):
                row_bg = c_fail_bg if elem_alt else white
                elem_alt = not elem_alt
                fail_tbl_data.append([
                    Paragraph(str(elem.get("ifc_file", "")), st_small),
                    Paragraph(str(elem.get("step_id", "")),   st_small),
                    Paragraph(str(elem.get("name", "")),      st_small),
                    Paragraph(str(elem.get("type", "")),      st_small),
                    Paragraph(str(elem.get("current_value", "")), st_small),
                ])
                fail_tbl_style.append(
                    ("BACKGROUND", (0, e_idx), (-1, e_idx), row_bg))

            fail_col_w = [body_w * 0.2, body_w * 0.1, body_w * 0.28,
                          body_w * 0.2, body_w * 0.22]
            fail_table = Table(fail_tbl_data, colWidths=fail_col_w, repeatRows=1)
            fail_table.setStyle(TableStyle(fail_tbl_style))
            story.append(fail_table)
        else:
            # Summary mode: just show count
            story.append(Paragraph(
                f"{len(elements)} elemento(s) con fallo", st_small))
        story.append(Spacer(1, 0.3 * cm))

    if not any_fail:
        story.append(Paragraph(
            "Sin elementos con fallos — todos los checks obligatorios han pasado.",
            _style("no_fail", fontSize=9, textColor=c_green,
                   fontName="Helvetica-Bold")))

    # ── 5. Page footer via doc template ──────────────────────────────────────
    def _add_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(c_gray)
        footer_text = (f"Generado por IFC HUB — {date_str}    "
                       f"Página {doc.page}")
        canvas.drawString(margin, 0.8 * cm, footer_text)
        canvas.restoreState()

    doc.build(story, onFirstPage=_add_footer, onLaterPages=_add_footer)
    buf.seek(0)
    return buf.read()
